"""
Pengujian generator DOCX multi-unit (Fase 3C).

Tiga unit dari SKKNI berbeda digabung ke satu dokumen:
  J.63OPR00.009.2 (2018-056), J.59MTM00.011.1 (2018-107), M.70MKT00.014.1 (2022-124)

Memeriksa isi, struktur XML numbering, pemisahan halaman, dan hasil rendering
(jarak nomor-teks diukur dari PDF, bukan hanya dari properti OOXML).
"""

import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docx.oxml.ns import qn
import pymupdf as fitz

from app.multi_docx_generator import generate_multi_docx
from app.docx_generator import LEVEL_INDENT_TWIPS
from app.section_parser import flatten_nodes, normalize_text

JSON_DIR = project_root / 'data' / 'json' / 'cross_skkni'
OUT_DIR = project_root / 'output'
DOCX_PATH = OUT_DIR / 'multi_unit_test.docx'
PDF_PATH = OUT_DIR / 'multi_unit_test.pdf'

SPEC = {'units': [
    {'kode': 'J.63OPR00.009.2', 'skkni': '2018-056'},
    {'kode': 'J.59MTM00.011.1', 'skkni': '2018-107'},
    {'kode': 'M.70MKT00.014.1', 'skkni': '2022-124'},
]}

lulus, gagal = [], []


def cek(nama, kondisi, detail=''):
    (lulus if kondisi else gagal).append(nama)
    print(f"  [{'PASS' if kondisi else 'FAIL'}] {nama}" + (f" — {detail}" if detail and not kondisi else ''))


def rapat(s):
    return re.sub(r'\s+', '', s or '')


def main():
    print("=" * 100)
    print("PENGUJIAN DOCX MULTI-UNIT — FASE 3C")
    print("=" * 100)

    info = generate_multi_docx(SPEC, str(DOCX_PATH))
    units = [json.loads((JSON_DIR / f"{u['kode']}.json").read_text(encoding='utf-8'))
             for u in SPEC['units']]

    print(f"\nDOCX: {DOCX_PATH.relative_to(project_root)} "
          f"({DOCX_PATH.stat().st_size / 1024:.0f} KB) | unit: {info['jumlah_unit']}")

    doc = Document(str(DOCX_PATH))
    paras = doc.paragraphs

    # ------------------------------------------------------------- urutan unit
    print("\n1. URUTAN & IDENTITAS UNIT")
    kode_paras = [p.text.strip() for p in paras if p.text.strip().startswith('KODE UNIT:')]
    cek('jumlah unit = 3', len(kode_paras) == 3, str(len(kode_paras)))
    urutan = [k.replace('KODE UNIT:', '').strip() for k in kode_paras]
    cek('urutan kode sesuai spesifikasi',
        urutan == [u['kode_unit'] for u in units], str(urutan))

    teks_doc = '\n'.join(p.text for p in paras)
    for u in units:
        cek(f"{u['kode_unit']}: judul & deskripsi ada",
            u['judul_unit'] in teks_doc and normalize_text(u['deskripsi_unit']) in normalize_text(teks_doc))

    # ------------------------------------------------------------ tabel & KUK
    print("\n2. TABEL, ELEMEN, DAN KUK")
    cek('3 tabel Word', len(doc.tables) == 3, str(len(doc.tables)))
    for tabel, u in zip(doc.tables, units):
        cek(f"{u['kode_unit']}: 2 kolom", len(tabel.columns) == 2)
        cek(f"{u['kode_unit']}: 1 header + {len(u['elemen_kompetensi'])} elemen",
            len(tabel.rows) == 1 + len(u['elemen_kompetensi']), str(len(tabel.rows)))
        kuk_docx = []
        for r in tabel.rows[1:]:
            kuk_docx.extend(p.text.split('\t', 1)[0].strip() for p in r.cells[1].paragraphs)
        kuk_json = [k['nomor_kuk'] for e in u['elemen_kompetensi'] for k in e['kuk']]
        cek(f"{u['kode_unit']}: nomor KUK lengkap & berurutan", kuk_docx == kuk_json,
            str([(a, b) for a, b in zip(kuk_docx, kuk_json) if a != b][:3]))

    # ------------------------------------------------ batasan & panduan lengkap
    print("\n3. KELENGKAPAN BATASAN VARIABEL & PANDUAN PENILAIAN")
    for u in units:
        nodes = (flatten_nodes(u['batasan_variabel']['sections'])
                 + flatten_nodes(u['panduan_penilaian']['sections']))
        hilang = [n['nomor'] for n in nodes if rapat(n['teks'])[:40] not in rapat(teks_doc)]
        cek(f"{u['kode_unit']}: seluruh teks bagian ada", not hilang, str(hilang[:5]))

    # Tidak ada duplikasi: jumlah kemunculan tiap teks = jumlah unit yang memuatnya
    from collections import Counter
    harap = Counter()
    for u in units:
        for n in (flatten_nodes(u['batasan_variabel']['sections'])
                  + flatten_nodes(u['panduan_penilaian']['sections'])):
            if len(n['teks']) > 25:
                harap[rapat(n['teks'])] += 1
        for e in u['elemen_kompetensi']:
            for k in e['kuk']:
                if len(k['uraian_kuk']) > 25:
                    harap[rapat(k['uraian_kuk'])] += 1
    teks_lengkap = teks_doc
    for tb in doc.tables:
        for r in tb.rows:
            for c in r.cells:
                teks_lengkap += '\n' + '\n'.join(p.text for p in c.paragraphs)
    teks_all = rapat(teks_lengkap)
    salah = [(t[:50], harap[t], teks_all.count(t)) for t in harap
             if teks_all.count(t) != harap[t]]
    cek('tidak ada teks terduplikasi', not salah, str(salah[:3]))

    # -------------------------------------------------------- numbering restart
    print("\n4. NUMBERING: KELOMPOK INDEPENDEN & RESTART")
    num_ids_semua = [nid for m in info['unit'] for nid in m['num_ids'].values()]
    cek('6 numId unik (2 per unit)', len(set(num_ids_semua)) == 6, str(num_ids_semua))

    ber_nomor = []
    for p in paras:
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn('w:numPr')) is not None:
            numPr = pPr.find(qn('w:numPr'))
            ber_nomor.append({
                'ilvl': int(numPr.find(qn('w:ilvl')).get(qn('w:val'))),
                'numId': int(numPr.find(qn('w:numId')).get(qn('w:val'))),
                'teks': p.text,
            })

    # Simulasi counter Word per numId
    counter = {nid: [0] * 4 for nid in set(num_ids_semua)}
    tampil = []
    for b in ber_nomor:
        nid, ilvl = b['numId'], b['ilvl']
        counter[nid][ilvl] += 1
        for d in range(ilvl + 1, 4):
            counter[nid][d] = 0
        tampil.append('.'.join(str(counter[nid][i]) for i in range(ilvl + 1)))

    idx = 0
    for u, m in zip(units, info['unit']):
        nodes = [n for n in (flatten_nodes(u['batasan_variabel']['sections'])
                             + flatten_nodes(u['panduan_penilaian']['sections']))
                 if n['is_numbered']]
        bagian = tampil[idx:idx + len(nodes)]
        idx += len(nodes)
        cek(f"{u['kode_unit']}: nomor tampilan = nomor JSON",
            bagian == [n['nomor'] for n in nodes],
            str([(a, b['nomor']) for a, b in zip(bagian, nodes) if a != b['nomor']][:3]))
        n_bv = len([n for n in flatten_nodes(u['batasan_variabel']['sections'])
                    if n['is_numbered']])
        cek(f"{u['kode_unit']}: BV & PP masing-masing mulai dari 1",
            bagian[0] == '1' and len(bagian) > n_bv and bagian[n_bv] == '1',
            f"BV={bagian[0]} PP={bagian[n_bv] if len(bagian) > n_bv else None}")

    cek('tidak ada numId dipakai dua unit',
        len({b['numId'] for b in ber_nomor}) == 6,
        str(sorted({b['numId'] for b in ber_nomor})))

    # ------------------------------------------------------- xml & page break
    print("\n5. STRUKTUR XML & PEMISAHAN HALAMAN")
    with zipfile.ZipFile(DOCX_PATH) as z:
        doc_xml = z.read('word/document.xml').decode('utf-8')
        num_xml = z.read('word/numbering.xml').decode('utf-8')

    cek('3 abstractNum untuk 3 unit', num_xml.count('<w:abstractNum ') >= 3)
    cek('setiap tabel punya tblHeader', doc_xml.count('<w:tblHeader/>') == 3,
        str(doc_xml.count('<w:tblHeader/>')))
    cek('semua baris cantSplit', doc_xml.count('<w:cantSplit/>') == sum(len(t.rows) for t in doc.tables),
        str(doc_xml.count('<w:cantSplit/>')))
    pb = [p for p in paras if p.paragraph_format.page_break_before]
    cek('2 page-break-before (unit 2 dan 3)', len(pb) == 2, str(len(pb)))
    cek('page break menempel di KODE UNIT unit berikutnya',
        all(p.text.strip().startswith('KODE UNIT:') for p in pb), str([p.text[:30] for p in pb]))

    # ------------------------------------------------------------- render PDF
    print("\n6. RENDER & PEMERIKSAAN VISUAL")
    dirender = False
    if shutil.which('flatpak'):
        subprocess.run(['flatpak', 'run', 'org.libreoffice.LibreOffice', '--headless',
                        '--convert-to', 'pdf', '--outdir', str(OUT_DIR), str(DOCX_PATH)],
                       capture_output=True, text=True, timeout=300)
        dirender = PDF_PATH.exists()

    if dirender:
        pdf = fitz.open(str(PDF_PATH))
        cek('PDF terbentuk', pdf.page_count > 0, str(pdf.page_count))

        halaman_unit = {}
        for i, page in enumerate(pdf):
            m = re.search(r'KODE UNIT: (\S+)', page.get_text())
            if m:
                halaman_unit[m.group(1)] = i + 1
                # unit harus mulai di bagian atas halaman
                y = next((l['bbox'][1] for b in page.get_text('dict')['blocks']
                          for l in b.get('lines', [])
                          if ''.join(s['text'] for s in l['spans']).strip().startswith('KODE UNIT:')), 999)
                cek(f"{m.group(1)}: mulai di halaman baru (y={y:.0f})", y < 120)

        cek('urutan halaman sesuai urutan unit',
            [halaman_unit.get(u['kode_unit']) for u in units]
            == sorted(halaman_unit.get(u['kode_unit']) for u in units),
            str(halaman_unit))

        # halaman kosong: tidak boleh ada halaman tanpa teks berarti
        kosong = [i + 1 for i, p in enumerate(pdf) if len(p.get_text().strip()) < 40]
        cek('tidak ada halaman kosong', not kosong, str(kosong))

        # jarak nomor-teks per level
        gap = {1: [], 2: [], 3: []}
        for page in pdf:
            for b in page.get_text('dict')['blocks']:
                for l in b.get('lines', []):
                    sp = l['spans']
                    if len(sp) < 3:
                        continue
                    nomor = sp[0]['text'].strip()
                    if not re.fullmatch(r'\d+(?:\.\d+){0,2}\.', nomor):
                        continue
                    level = nomor.count('.')
                    teks_x = sp[-1]['bbox'][0]
                    gap.setdefault(level, []).append(teks_x - sp[0]['bbox'][2])
        for level, nilai in sorted(gap.items()):
            if not nilai:
                continue
            cek(f'level {level}: jarak nomor-teks wajar (4-12 pt)',
                all(4 <= x <= 12 for x in nilai),
                f"min={min(nilai):.1f} max={max(nilai):.1f}")
            cek(f'level {level}: jarak konsisten (selisih <= 2 pt)',
                max(nilai) - min(nilai) <= 2,
                f"min={min(nilai):.1f} max={max(nilai):.1f}")

        # teks lengkap di PDF
        teks_pdf = rapat('\n'.join(p.get_text() for p in pdf))
        hilang_pdf = []
        for u in units:
            for n in (flatten_nodes(u['batasan_variabel']['sections'])
                      + flatten_nodes(u['panduan_penilaian']['sections'])):
                if rapat(n['teks'])[:40] not in teks_pdf:
                    hilang_pdf.append(n['nomor'])
        cek('seluruh teks bagian tampil di PDF', not hilang_pdf, str(hilang_pdf[:5]))
        print(f"     halaman PDF: {pdf.page_count} | unit mulai: {halaman_unit}")
        pdf.close()
    else:
        print('  [SKIP] LibreOffice tidak tersedia')

    # ------------------------------------------------------------- regresi
    print("\n7. REGRESI")
    for nama in ('test_extractor.py', 'test_unit_parser.py', 'test_unit_complete.py',
                 'test_unit_numbering.py', 'test_docx_generator.py'):
        r = subprocess.run([sys.executable, str(project_root / 'tests' / nama)],
                           capture_output=True, text=True, cwd=str(project_root))
        cek(f'{nama} lulus', r.returncode == 0,
            (r.stdout + r.stderr).strip().splitlines()[-1] if r.returncode else '')

    # ------------------------------------------------------------ ringkasan
    print("\n" + "-" * 100)
    print(f"Unit           : {info['jumlah_unit']} ({', '.join(info['urutan_kode'])})")
    print(f"Total Elemen   : {info['total_elemen']}")
    print(f"Total KUK      : {info['total_kuk']}")
    print("\n" + "=" * 100)
    print(f"HASIL: {len(lulus)} lulus, {len(gagal)} gagal")
    if gagal:
        print("GAGAL:")
        for g in gagal:
            print(f"  - {g}")
    print("=" * 100)
    return 0 if not gagal else 1




def test_multi_docx():
    assert main() == 0


if __name__ == '__main__':
    sys.exit(main())
