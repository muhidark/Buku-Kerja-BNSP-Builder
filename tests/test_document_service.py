"""
Pengujian layanan pembuatan dokumen Word (Fase 4E).

DOCX dan metadata ditulis ke direktori uji agar riwayat produksi tidak
tercampur. Sumber JSON dan skema hanya dibaca.
"""

import json
import re
import shutil
import sys
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docx.oxml.ns import qn

from app import jobs, schema_library as sl
from app.services import document_service as ds
from app.section_parser import flatten_nodes

DOC_UJI = project_root / 'output' / 'web_documents' / '_uji'
META_UJI = project_root / 'data' / 'web_documents' / '_uji'

lulus, gagal = [], []


def cek(nama, kondisi, detail=''):
    (lulus if kondisi else gagal).append(nama)
    print(f"  [{'PASS' if kondisi else 'FAIL'}] {nama}" + (f" — {detail}" if detail and not kondisi else ''))


def pasang_uji():
    for d in (DOC_UJI, META_UJI):
        if d.exists():
            shutil.rmtree(d)
        d.mkdir(parents=True, exist_ok=True)
    ds.DOC_DIR = DOC_UJI
    ds.META_DIR = META_UJI


def bersihkan():
    from app import schema_library as _sl
    if getattr(main, '_asli_ambil', None) is not None:
        _sl.ambil_skema = main._asli_ambil
    ds.DOC_DIR = project_root / 'output' / 'web_documents'
    ds.META_DIR = project_root / 'data' / 'web_documents'
    for d in (DOC_UJI, META_UJI):
        if d.exists():
            shutil.rmtree(d, ignore_errors=True)


def teks_docx(path: Path) -> str:
    doc = Document(str(path))
    bagian = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                bagian.extend(p.text for p in c.paragraphs)
    return '\n'.join(bagian)


def main():
    print("=" * 100)
    print("PENGUJIAN DOKUMEN WORD — FASE 4E (SERVICE)")
    print("=" * 100)

    sid = next((d['id'] for d in sl.daftar_skema()
                if d['nama_skema'] == 'SOCIAL MEDIA MARKETING'), None)
    if not sid:
        print('ERROR: skema SOCIAL MEDIA MARKETING tidak ada.')
        return 1

    pasang_uji()

    # Skema produksi dapat memiliki penandaan manual dari pengguna. Untuk
    # baseline pengujian, penandaan itu dilepas dulu (lihat bagian 9).
    asli_ambil_skema = sl.ambil_skema
    main._asli_ambil = asli_ambil_skema

    def tanpa_manual(_id):
        asli = asli_ambil_skema(_id)
        if not asli:
            return asli
        d = dict(asli)
        d['verifikasi_manual'] = {}
        return d

    sl.ambil_skema = tanpa_manual

    skema = sl.ambil_skema(sid)
    kode_skema = [u['kode_unit'] for u in skema['units']]
    print(f"\nSkema: {skema['nama_skema']} | unit: {len(kode_skema)} | {kode_skema}")

    # --------------------------------------------------------- 1 validasi
    print("\n1. VALIDASI SEBELUM GENERATE")
    v = ds.validasi(sid)
    cek('empat unit diperiksa', len(v['unit']) == 4, str(len(v['unit'])))
    cek('urutan sesuai skema', [u['kode_unit'] for u in v['unit']] == kode_skema)
    cek('ada unit perlu verifikasi', v['lengkap'] is False)
    cek('masalah menunjuk unit yang tepat',
        any(m['jenis'] == 'validasi_struktur' for m in v['masalah']),
        str(v['masalah']))
    cek('semua unit punya JSON', all(u['json'] for u in v['unit']))

    # ------------------------------------------------------ 2 draft ditolak
    print("\n2. DRAFT TANPA PERSETUJUAN DITOLAK")
    try:
        ds.buat_dokumen(sid)
        cek('penolakan draft tanpa persetujuan', False, 'tidak menolak')
    except ds.GalatDokumen as e:
        cek('penolakan draft tanpa persetujuan', e.kode == 409, str(e.kode))
        cek('alasan penolakan disertakan', bool(e.masalah))
    cek('tidak ada berkas dibuat saat ditolak', not list(META_UJI.glob('*.json')))

    # ------------------------------------------------- 3 draft disetujui
    print("\n3. DRAFT DENGAN PERSETUJUAN")
    meta = ds.buat_dokumen(sid, setujui_draft=True)
    cek('status draft_perlu_verifikasi', meta['status'] == 'draft_perlu_verifikasi', meta['status'])
    cek('persetujuan draft tercatat', meta['persetujuan_draft'] is True)
    cek('empat unit masuk', meta['jumlah_unit'] == 4, str(meta['jumlah_unit']))
    cek('urutan sesuai skema', meta['urutan_kode'] == kode_skema, str(meta['urutan_kode']))
    path = project_root / meta['berkas_docx']
    cek('berkas DOCX benar-benar ada', path.exists() and path.stat().st_size > 10000,
        str(path.stat().st_size if path.exists() else 0))

    # --------------------------------------------------- 4 isi dokumen
    print("\n4. ISI DOKUMEN")
    doc = Document(str(path))
    cek('jumlah tabel = jumlah unit', len(doc.tables) == 4, str(len(doc.tables)))
    kode_doc = [p.text.split(':', 1)[1].strip() for p in doc.paragraphs
                if p.text.strip().startswith('KODE UNIT:')]
    cek('kode unit di DOCX sesuai urutan', kode_doc == kode_skema, str(kode_doc))

    total_kuk_doc = sum(len(r.cells[1].paragraphs) for t in doc.tables for r in t.rows[1:])
    cek('total KUK sama dengan metadata', total_kuk_doc == meta['total_kuk'],
        f"{total_kuk_doc} vs {meta['total_kuk']}")

    # numbering: setiap unit punya numId sendiri
    num_ids = set()
    for p in doc.paragraphs:
        numPr = p._p.pPr.find(qn('w:numPr')) if p._p.pPr is not None else None
        if numPr is not None:
            num_ids.add(int(numPr.find(qn('w:numId')).get(qn('w:val'))))
    cek('numbering per unit independen', len(num_ids) == 8, str(sorted(num_ids)))

    # seluruh teks BV/PP ada
    teks = re.sub(r'\s+', ' ', teks_docx(path))
    hilang = []
    for kode in kode_skema:
        data = json.loads((project_root / 'data' / 'json' / 'browser_extraction'
                           / f'{kode}.json').read_text(encoding='utf-8'))
        for n in (flatten_nodes(data['batasan_variabel']['sections'])
                  + flatten_nodes(data['panduan_penilaian']['sections'])):
            if re.sub(r'\s+', '', n['teks'])[:50] not in re.sub(r'\s+', '', teks):
                hilang.append((kode, n['nomor']))
    cek('seluruh teks BV/PP masuk DOCX', not hilang, str(hilang[:3]))

    # ------------------------------------------------- 5 unit terpilih
    print("\n5. GENERATE UNIT TERPILIH")
    dua = [kode_skema[0], kode_skema[2]]
    m2 = ds.buat_dokumen(sid, list(reversed(dua)), setujui_draft=True)
    cek('dua unit masuk', m2['jumlah_unit'] == 2)
    cek('urutan dipulihkan sesuai skema',
        m2['urutan_kode'] == [k for k in kode_skema if k in dua], str(m2['urutan_kode']))
    cek('status siap_ditinjau untuk unit yang lolos',
        m2['status'] == 'siap_ditinjau', m2['status'])

    satu = [kode_skema[1]]
    try:
        ds.buat_dokumen(sid, satu)
        cek('unit perlu verifikasi ditolak tanpa persetujuan', False, 'tidak menolak')
    except ds.GalatDokumen as e:
        cek('unit perlu verifikasi ditolak tanpa persetujuan', e.kode == 409, str(e.kode))

    m1 = ds.buat_dokumen(sid, satu, setujui_draft=True)
    cek('satu unit perlu verifikasi -> draft', m1['status'] == 'draft_perlu_verifikasi',
        m1['status'])
    cek('satu unit masuk', m1['jumlah_unit'] == 1)

    # -------------------------------------------------- 6 penolakan input
    print("\n6. PENOLAKAN INPUT")
    for nama, fn in [
        ('kode duplikat', lambda: ds.buat_dokumen(sid, [kode_skema[0], kode_skema[0]], True)),
        ('kode asing', lambda: ds.buat_dokumen(sid, ['X.00XXX00.000.0'], True)),
        ('skema tidak ada', lambda: ds.buat_dokumen('0' * 32, None, True)),
    ]:
        try:
            fn()
            cek(f'{nama} ditolak', False, 'tidak menolak')
        except ds.GalatDokumen as e:
            cek(f'{nama} ditolak', e.kode in (400, 404), f'{e.kode}: {e.pesan[:40]}')

    # ------------------------------------------------ 7 JSON tidak tersedia
    print("\n7. JSON TIDAK TERSEDIA")
    asli_cari = ds._cari_json
    ds._cari_json = lambda kode: (None if kode == kode_skema[0] else asli_cari(kode))
    try:
        v2 = ds.validasi(sid)
        cek('kekurangan JSON terdeteksi',
            any(m['jenis'] == 'json_tidak_ada' for m in v2['masalah']), str(v2['masalah']))
        try:
            ds.buat_dokumen(sid, [kode_skema[0]], setujui_draft=True)
            cek('unit tanpa JSON tidak dapat dibuat', False)
        except ds.GalatDokumen as e:
            cek('unit tanpa JSON tidak dapat dibuat', e.kode == 409, str(e.kode))
    finally:
        ds._cari_json = asli_cari

    # ------------------------------------------------------- 8 riwayat
    print("\n8. RIWAYAT & METADATA")
    daftar = ds.daftar_dokumen()
    cek('riwayat memuat dokumen yang dibuat', len(daftar) >= 3, str(len(daftar)))
    cek('metadata atomik tersimpan',
        all((META_UJI / f"{d['id']}.json").exists() for d in daftar))
    isi = json.loads((META_UJI / f"{meta['id']}.json").read_text(encoding='utf-8'))
    wajib = ('id', 'schema_id', 'units', 'berkas_docx', 'status', 'dibuat', 'total_kuk')
    cek('kunci metadata lengkap', all(k in isi for k in wajib),
        str([k for k in wajib if k not in isi]))
    cek('metadata mencatat hash sumber',
        all(u.get('sha256_pdf') for u in isi['units']))
    cek('tidak menampilkan path absolut',
        not any(str(project_root) in json.dumps(d) for d in daftar))

    print("\n9. VERIFIKASI MANUAL MENYELESAIKAN TEMUAN")
    sl.ambil_skema = lambda _id: {**tanpa_manual(_id),
                                  'verifikasi_manual': {kode_skema[1]: {'waktu': 'uji'}}}
    try:
        v3 = ds.validasi(sid)
        cek('tanpa manual: belum lengkap', v['lengkap'] is False)
        cek('dengan manual: lengkap', v3['lengkap'] is True, str(v3['masalah']))
        cek('unit ditandai manual', any(u.get('terverifikasi_manual') for u in v3['unit']))
        m3 = ds.buat_dokumen(sid)
        cek('dokumen final dapat dibuat', m3['status'] == 'siap_ditinjau', m3['status'])
        cek('metadata mencatat verifikasi manual',
            any(u['terverifikasi_manual'] for u in m3['units']))
    finally:
        sl.ambil_skema = tanpa_manual

    print("\n10. STATUS EKSTRAKSI TIDAK BERUBAH")
    data_awal = json.loads((project_root / 'data' / 'json' / 'browser_extraction'
                            / f'{kode_skema[0]}.json').read_text(encoding='utf-8'))
    cek('JSON sumber tidak diubah oleh generator',
        data_awal['kode_unit'] == kode_skema[0]
        and (data_awal.get('validasi') or {}).get('status') == 'valid_struktur')

    bersihkan()
    print("\n" + "=" * 100)
    print(f"HASIL: {len(lulus)} lulus, {len(gagal)} gagal")
    if gagal:
        for g in gagal:
            print(f"  - {g}")
    print("=" * 100)
    return 0 if not gagal else 1




def test_document_service():
    assert main() == 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    finally:
        bersihkan()
