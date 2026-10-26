"""
Pengujian lintas SKKNI (Fase 3B).

Alur: PDF -> identifikasi unit -> JSON -> DOCX -> PDF (render) -> validasi.

Empat unit diuji:
  J.63OPR00.009.2  SKKNI 2018-056
  J.59MTM00.011.1  SKKNI 2018-107
  M.70MKT00.014.1  SKKNI 2022-124
  M.73ADV00.001.1  SKKNI 2023-101

Output:
  data/json/cross_skkni/<kode>.json
  output/cross_skkni/<kode>.docx
  output/cross_skkni/<kode>.pdf
  output/cross_skkni/laporan_pengujian.json
"""

import json
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import datetime
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docx.oxml.ns import qn
import pymupdf as fitz

from app.extractor import identify_units
from app.unit_parser import parse_unit, save_unit_json
from app.docx_generator import (
    generate_docx, LEVEL_INDENT_TWIPS, HANGING_TWIPS, FONT_NAME, FONT_SIZE_PT,
)
from app.section_parser import flatten_nodes, normalize_text

PDF_DIR = project_root / 'data' / 'pdf'
JSON_DIR = project_root / 'data' / 'json' / 'cross_skkni'
OUT_DIR = project_root / 'output' / 'cross_skkni'
REPORT = OUT_DIR / 'laporan_pengujian.json'

JOBS = [
    {'skkni': 'SKKNI 2018-056', 'pdf': 'SKKNI 2018-056.pdf.pdf',
     'kode': 'J.63OPR00.009.2',
     'judul': 'Menggunakan Aplikasi Berbasis Internet (Internet Based Applications Literacy)',
     # Regresi Fase 3B.2: satu subbutir harus jadi satu paragraf utuh
     'pp_2_1': ('J.63OPR00.007.2 Menggunakan Penelusur Situs Web (Web Browser)')},
    {'skkni': 'SKKNI 2018-107', 'pdf': 'SKKNI 2018-107.pdf.pdf',
     'kode': 'J.59MTM00.011.1',
     'judul': 'Membuat Aset Visual Berdasarkan Langkah Kerja yang Telah Ditetapkan',
     # Acuan audit Fase 3B: deskripsi tidak boleh menyerap Batasan Variabel
     'deskripsi': ('Unit ini berhubungan dengan pengetahuan, keterampilan dan sikap kerja '
                   'yang dibutuhkan untuk dapat membuat aset visual berdasarkan langkah '
                   'kerja yang telah ditetapkan.')},
    {'skkni': 'SKKNI 2022-124', 'pdf': 'SKKNI 2022-124.pdf.pdf',
     'kode': 'M.70MKT00.014.1',
     'judul': 'Mempersiapkan Konten Digital'},
    {'skkni': 'SKKNI 2023-101', 'pdf': 'SKKNI 2023-101.pdf.pdf',
     'kode': 'M.73ADV00.001.1',
     'judul': 'Menyusun Rencana Pengembangan Usaha Perusahaan Periklanan'},
]

lulus, gagal = [], []


def cek(unit, nama, kondisi, detail=''):
    (lulus if kondisi else gagal).append((unit, nama))
    print(f"  [{'PASS' if kondisi else 'FAIL'}] {nama}" + (f" — {detail}" if detail and not kondisi else ''))


def rapat(s: str) -> str:
    return re.sub(r'\s+', '', s or '')


def indent_harap(level: int) -> int:
    return LEVEL_INDENT_TWIPS.get(level, LEVEL_INDENT_TWIPS[3] + (level - 3) * HANGING_TWIPS)


def render_pdf(docx_path: Path) -> bool:
    if not shutil.which('flatpak'):
        return False
    subprocess.run(['flatpak', 'run', 'org.libreoffice.LibreOffice', '--headless',
                    '--convert-to', 'pdf', '--outdir', str(OUT_DIR), str(docx_path)],
                   capture_output=True, text=True, timeout=300)
    return docx_path.with_suffix('.pdf').exists()


def validasi_json(unit: dict, job: dict) -> dict:
    hasil = {'catatan': [], 'perlu_verifikasi': 0}
    kode = unit['kode_unit']

    cek(kode, 'kode unit sesuai', kode == job['kode'], kode)
    cek(kode, 'judul sesuai PDF/target',
        rapat(unit['judul_unit']).lower().startswith(rapat(job['judul']).lower()[:40]),
        unit['judul_unit'])
    cek(kode, 'deskripsi terisi', bool(unit['deskripsi_unit']))

    # Identitas tidak boleh mengandung label struktur lain (kasus J.59MTM00.011.1)
    label_struktur = ('BATASAN VARIABEL', 'PANDUAN PENILAIAN', 'ELEMEN KOMPETENSI',
                      'KRITERIA UNJUK KERJA', 'KODE UNIT', 'JUDUL UNIT')
    cemar = [l for l in label_struktur
             if l in (unit['deskripsi_unit'] or '').upper()
             or l in (unit['judul_unit'] or '').upper()]
    cek(kode, 'identitas bebas label struktur', not cemar, str(cemar))
    cek(kode, 'deskripsi tidak memuat teks Batasan/Panduan',
        not any(x in (unit['deskripsi_unit'] or '').upper()
                for x in ('BATASAN VARIABEL', 'PANDUAN PENILAIAN')))

    if job.get('deskripsi'):
        cek(kode, 'deskripsi persis sesuai PDF',
            normalize_text(unit['deskripsi_unit']) == normalize_text(job['deskripsi']),
            repr((unit['deskripsi_unit'] or '')[:100]))

    # Satu subbutir = satu teks utuh (tidak pecah per kata)
    if job.get('pp_2_1'):
        pp = flatten_nodes(unit['panduan_penilaian']['sections'])
        node = next((n for n in pp if n['nomor'] == '2.1'), None)
        cek(kode, 'PP 2.1 menjadi satu paragraf utuh',
            node is not None and normalize_text(node['teks']) == normalize_text(job['pp_2_1']),
            repr(node['teks'] if node else None))

    # Peringatan format sumber tidak boleh menjadikan status gagal
    peringatan = unit['validasi'].get('peringatan_format_sumber', [])
    cek(kode, 'peringatan berkategori peringatan_format_sumber',
        all(x.get('kategori') == 'peringatan_format_sumber' for x in peringatan),
        str(peringatan[:2]))
    cek(kode, 'peringatan tidak menggagalkan status',
        not peringatan or unit['validasi']['status'] == 'valid_struktur',
        f"{unit['validasi']['status']} / {len(peringatan)} peringatan")
    cek(kode, 'KUK tanpa titik tidak diberi titik tambahan',
        all(not k['uraian_kuk'].endswith('.') for e in unit['elemen_kompetensi']
            for k in e['kuk'] if k['nomor_kuk'] in {x['nomor_kuk'] for x in peringatan}))

    elemen = unit['elemen_kompetensi']
    cek(kode, 'elemen berurutan 1..n',
        [e['nomor_elemen'] for e in elemen] == [str(i) for i in range(1, len(elemen) + 1)],
        str([e['nomor_elemen'] for e in elemen]))

    for e in elemen:
        harap = [f"{e['nomor_elemen']}.{i}" for i in range(1, len(e['kuk']) + 1)]
        cek(kode, f"KUK elemen {e['nomor_elemen']} berurutan",
            [k['nomor_kuk'] for k in e['kuk']] == harap,
            str([k['nomor_kuk'] for k in e['kuk']]))

    for bagian in ('batasan_variabel', 'panduan_penilaian'):
        b = unit[bagian]
        nodes = flatten_nodes(b['sections'])
        cek(kode, f'{bagian} ditemukan & berisi', b['ditemukan'] and bool(nodes))
        cek(kode, f'{bagian} penomoran mulai dari 1',
            bool(b['penomoran']['nomor_akar']) and b['penomoran']['nomor_akar'][0] == '1',
            str(b['penomoran']['nomor_akar'][:3]))

    if unit['validasi']['status'] != 'valid_struktur':
        hasil['perlu_verifikasi'] += 1
        hasil['catatan'].append(f"status: {unit['validasi']['status']}")
        hasil['catatan'].extend(unit['validasi']['catatan'])

    level_besar = (unit['batasan_variabel']['penomoran']['level_melebihi_tiga']
                   + unit['panduan_penilaian']['penomoran']['level_melebihi_tiga'])
    if level_besar:
        hasil['perlu_verifikasi'] += 1
        hasil['catatan'].append(f'level di atas 3: {sorted(set(level_besar))} (perlu penentuan format)')

    return hasil


def validasi_pdf_ke_json(pdf_path: str, unit: dict) -> dict:
    doc = fitz.open(pdf_path)
    kode = unit['kode_unit']

    def teks_hal(a, b):
        b = b if b is not None else a
        return normalize_text(' '.join(doc[p].get_text() for p in range(a, b + 1)))

    hasil = {'dicek': 0, 'cocok': 0, 'tidak_cocok': []}

    nodes = (flatten_nodes(unit['batasan_variabel']['sections'])
             + flatten_nodes(unit['panduan_penilaian']['sections']))
    for n in nodes:
        t = teks_hal(n['halaman_pdf'], n['halaman_akhir'])
        for baris in (n['teks_mentah'] or '').split('\n'):
            b = normalize_text(baris)
            if not b:
                continue
            hasil['dicek'] += 1
            if b in t:
                hasil['cocok'] += 1
            else:
                hasil['tidak_cocok'].append({'nomor': n['nomor'], 'baris': b[:80]})

    for e in unit['elemen_kompetensi']:
        for kk in e['kuk']:
            t = teks_hal(kk['halaman_pdf'], kk['halaman_pdf'])
            for baris in (kk['teks_mentah'] or '').split('\n'):
                b = normalize_text(baris)
                if not b:
                    continue
                hasil['dicek'] += 1
                if b in t:
                    hasil['cocok'] += 1
                else:
                    hasil['tidak_cocok'].append({'nomor': kk['nomor_kuk'], 'baris': b[:80]})

    doc.close()
    cek(kode, 'PDF->JSON: seluruh baris teks ditemukan di PDF',
        not hasil['tidak_cocok'],
        f"{hasil['cocok']}/{hasil['dicek']}; contoh {hasil['tidak_cocok'][:2]}")
    return hasil


def validasi_json_ke_docx(docx_path: Path, unit: dict, info: dict) -> dict:
    kode = unit['kode_unit']
    doc = Document(str(docx_path))
    hasil = {'catatan': []}

    bv = flatten_nodes(unit['batasan_variabel']['sections'])
    pp = flatten_nodes(unit['panduan_penilaian']['sections'])
    semua = bv + pp
    bernomor = [n for n in semua if n['is_numbered']]
    tanpa = [n for n in semua if not n['is_numbered']]
    kuk_json = [k for e in unit['elemen_kompetensi'] for k in e['kuk']]

    # Tabel
    cek(kode, 'DOCX: 1 tabel Word', len(doc.tables) == 1, str(len(doc.tables)))
    tabel = doc.tables[0]
    cek(kode, 'DOCX: 2 kolom', len(tabel.columns) == 2)
    cek(kode, 'DOCX: 1 header + n elemen',
        len(tabel.rows) == 1 + len(unit['elemen_kompetensi']), str(len(tabel.rows)))
    cek(kode, 'DOCX: header tabel tebal & benar',
        tabel.rows[0].cells[0].text.strip() == 'ELEMEN KOMPETENSI'
        and all(r.bold for c in tabel.rows[0].cells for p in c.paragraphs for r in p.runs))

    kuk_docx = []
    for i, e in enumerate(unit['elemen_kompetensi'], start=1):
        sel = tabel.rows[i].cells[1]
        kuk_docx.extend(sel.paragraphs)
    nomor_docx = [p.text.split('\t', 1)[0].strip() for p in kuk_docx]
    cek(kode, 'DOCX: jumlah & nomor KUK sesuai JSON',
        nomor_docx == [k['nomor_kuk'] for k in kuk_json],
        str([(a, b) for a, b in zip(nomor_docx, [k['nomor_kuk'] for k in kuk_json]) if a != b][:3]))
    cek(kode, 'DOCX: uraian KUK identik',
        all(p.text.split('\t', 1)[-1].strip() == k['uraian_kuk'] for p, k in zip(kuk_docx, kuk_json)))

    # Paragraf bernomor
    ber_nomor = []
    for p in doc.paragraphs:
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn('w:numPr')) is not None:
            numPr = pPr.find(qn('w:numPr'))
            ind = pPr.find(qn('w:ind'))
            ber_nomor.append({
                'ilvl': int(numPr.find(qn('w:ilvl')).get(qn('w:val'))),
                'numId': int(numPr.find(qn('w:numId')).get(qn('w:val'))),
                'left': int(ind.get(qn('w:left'))) if ind is not None and ind.get(qn('w:left')) else None,
                'hanging': int(ind.get(qn('w:hanging'))) if ind is not None and ind.get(qn('w:hanging')) else 0,
            })
    cek(kode, 'DOCX: jumlah paragraf bernomor sesuai JSON',
        len(ber_nomor) == len(bernomor), f'{len(ber_nomor)} vs {len(bernomor)}')

    # Simulasi nomor tampilan
    counter = {nid: [0] * info['levels'] for nid in info['num_ids'].values()}
    tampil = []
    for b in ber_nomor:
        if b['numId'] not in counter:
            tampil.append(None)
            continue
        ilvl = b['ilvl']
        counter[b['numId']][ilvl] += 1
        for d in range(ilvl + 1, info['levels']):
            counter[b['numId']][d] = 0
        tampil.append('.'.join(str(counter[b['numId']][i]) for i in range(ilvl + 1)))
    cek(kode, 'DOCX: nomor tampilan = nomor JSON',
        tampil == [n['nomor'] for n in bernomor],
        str([(a, b['nomor']) for a, b in zip(tampil, bernomor) if a != b['nomor']][:3]))

    # Indentasi
    salah = [(b['ilvl'], b['left'], b['hanging']) for b in ber_nomor
             if b['left'] != indent_harap(b['ilvl'] + 1)
             or b['hanging'] != info['hanging_per_level'].get(b['ilvl'] + 1, HANGING_TWIPS)]
    cek(kode, 'DOCX: indentasi sesuai level', not salah, str(salah[:3]))

    # Paragraf tanpa nomor
    paragraf_tanpa = [p for p in doc.paragraphs
                      if (p._p.pPr is None or p._p.pPr.find(qn('w:numPr')) is None)
                      and p.text.strip() in {n['teks'] for n in tanpa}]
    cek(kode, 'DOCX: paragraf tanpa nomor tidak bernomor',
        len(paragraf_tanpa) == len(tanpa), f'{len(paragraf_tanpa)} vs {len(tanpa)}')

    # Heading bagian muncul tepat satu kali
    for heading in ('BATASAN VARIABEL', 'PANDUAN PENILAIAN'):
        jumlah = sum(1 for p in doc.paragraphs if p.text.strip() == heading)
        cek(kode, f'DOCX: heading {heading} muncul sekali', jumlah == 1, str(jumlah))

    # Deskripsi utuh dan tidak tercampur
    paragraf_desc = [p.text for p in doc.paragraphs
                     if normalize_text(p.text) == normalize_text(unit['deskripsi_unit'])]
    cek(kode, 'DOCX: deskripsi utuh tanpa teks bagian lain', len(paragraf_desc) == 1,
        str(len(paragraf_desc)))

    # Tidak ada paragraf fragmen: paragraf pendek yang bukan teks resmi JSON
    teks_resmi = {normalize_text(n['teks']) for n in semua}
    teks_resmi |= {normalize_text(unit['judul_unit']), normalize_text(unit['deskripsi_unit']),
                   'BATASAN VARIABEL', 'PANDUAN PENILAIAN'}
    fragmen = [p.text for p in doc.paragraphs
               if p.text.strip() and len(p.text.split()) <= 2
               and normalize_text(p.text) not in teks_resmi]
    cek(kode, 'DOCX: tidak ada paragraf fragmen', not fragmen, str(fragmen[:5]))

    # Kelengkapan teks
    teks_doc = '\n'.join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                teks_doc += '\n' + '\n'.join(p.text for p in c.paragraphs)
    teks_rapat = rapat(teks_doc)
    hilang = [n['teks'][:50] for n in semua if rapat(n['teks']) not in teks_rapat]
    hilang += [k['uraian_kuk'][:50] for k in kuk_json if rapat(k['uraian_kuk']) not in teks_rapat]
    cek(kode, 'DOCX: seluruh teks JSON ada', not hilang, str(hilang[:3]))

    # XML
    with zipfile.ZipFile(docx_path) as z:
        numbering_xml = z.read('word/numbering.xml').decode('utf-8')
        document_xml = z.read('word/document.xml').decode('utf-8')
    cek(kode, 'DOCX: abstractNum punya level sesuai data',
        numbering_xml.count(f'<w:lvl w:ilvl=') >= info['levels'])
    cek(kode, 'DOCX: dua numId dengan startOverride',
        all(re.search(rf'<w:num w:numId="{nid}".*?</w:num>', numbering_xml, re.S)
            and re.search(rf'<w:num w:numId="{nid}".*?</w:num>', numbering_xml, re.S)
            .group(0).count('<w:startOverride w:val="1"/>') == info['levels']
            for nid in info['num_ids'].values()))
    cek(kode, 'DOCX: header tabel diulang (tblHeader)', '<w:tblHeader/>' in document_xml)
    cek(kode, 'DOCX: semua baris cantSplit',
        document_xml.count('<w:cantSplit/>') == len(tabel.rows),
        f"{document_xml.count('<w:cantSplit/>')} vs {len(tabel.rows)}")
    cek(kode, 'DOCX: font & ukuran sesuai referensi',
        doc.styles['Normal'].font.name == FONT_NAME
        and doc.styles['Normal'].font.size.pt == FONT_SIZE_PT)

    hasil['jumlah_paragraf_bernomor'] = len(ber_nomor)
    hasil['jumlah_paragraf_tanpa_nomor'] = len(paragraf_tanpa)
    hasil['jumlah_tabel'] = len(doc.tables)
    return hasil


def validasi_docx_pdf(pdf_path: Path, unit: dict) -> dict:
    kode = unit['kode_unit']
    hasil = {'halaman': 0}
    if not pdf_path.exists():
        cek(kode, 'PDF render tersedia', False)
        return hasil

    pdf = fitz.open(str(pdf_path))
    teks = '\n'.join(p.get_text() for p in pdf)
    hasil['halaman'] = pdf.page_count
    cek(kode, 'PDF render terbentuk', pdf.page_count > 0, str(pdf.page_count))

    nodes = (flatten_nodes(unit['batasan_variabel']['sections'])
             + flatten_nodes(unit['panduan_penilaian']['sections']))
    bernomor = [n for n in nodes if n['is_numbered']]
    tanpa = [n for n in nodes if not n['is_numbered']]

    # Nomor terdalam harus tampil (bukan hanya properti numPr)
    sampel = [n['nomor'] for n in bernomor if n['level'] >= 2][:6]
    hilang = [n for n in sampel if f'{n}.' not in teks]
    cek(kode, 'PDF: penomoran tampil', not hilang, str(hilang[:3]))

    teks_rapat = rapat(teks)
    hilang_teks = [n['nomor'] for n in bernomor if rapat(n['teks'])[:40] not in teks_rapat]
    cek(kode, 'PDF: seluruh teks bagian tampil', not hilang_teks, str(hilang_teks[:3]))

    tanpa_ta = [n for n in tanpa if n['teks'].strip() == '(Tidak ada.)']
    jumlah_ta = len(re.findall(r'(?m)^\(Tidak ada\.\)', teks))
    cek(kode, 'PDF: "(Tidak ada.)" tanpa nomor', jumlah_ta == len(tanpa_ta),
        f'{jumlah_ta} vs {len(tanpa_ta)}')

    pdf.close()
    return hasil


def main():
    print("=" * 100)
    print("PENGUJIAN LINTAS SKKNI — FASE 3B")
    print("=" * 100)

    JSON_DIR.mkdir(parents=True, exist_ok=True)
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    baris_laporan = []
    perbedaan_teks = []
    peringatan_sumber = []
    masalah_format = []
    hasil_unit = []

    for job in JOBS:
        pdf_path = PDF_DIR / job['pdf']
        if not pdf_path.exists():
            cek(job['kode'], f'PDF {job["pdf"]} tersedia', False)
            continue

        print(f"\n{'=' * 100}\n{job['skkni']} — {job['kode']}\n{'=' * 100}")

        # Identifikasi unit (Fase 1) untuk rentang halaman
        units = identify_units(str(pdf_path))
        target = next((u for u in units if u['code'] == job['kode']), None)
        cek(job['kode'], 'unit ditemukan oleh Fase 1', target is not None)
        if target is None:
            continue

        unit = parse_unit(str(pdf_path), job['kode'],
                          target['start_page'], target['end_page'])
        unit['metadata_sumber']['nama_dokumen'] = job['pdf']
        item_json = JSON_DIR / f"{job['kode']}.json"
        save_unit_json(unit, str(item_json))

        print(f"  halaman PDF {target['start_page']}-{target['end_page']} | "
              f"elemen {len(unit['elemen_kompetensi'])} | "
              f"KUK {sum(len(e['kuk']) for e in unit['elemen_kompetensi'])} | "
              f"status {unit['validasi']['status']}")

        v_json = validasi_json(unit, job)
        v_pdf2json = validasi_pdf_ke_json(str(pdf_path), unit)

        # DOCX
        docx_path = OUT_DIR / f"{job['kode']}.docx"
        info = generate_docx(str(item_json), str(docx_path))
        cek(job['kode'], 'DOCX terbentuk', docx_path.exists() and docx_path.stat().st_size > 5000)
        v_json2docx = validasi_json_ke_docx(docx_path, unit, info)

        pdf_out = OUT_DIR / f"{job['kode']}.pdf"
        dirender = render_pdf(docx_path)
        v_docx2pdf = validasi_docx_pdf(pdf_out, unit) if dirender else {'halaman': 0}

        if v_pdf2json['tidak_cocok']:
            perbedaan_teks.append({
                'kode_unit': job['kode'],
                'jumlah': len(v_pdf2json['tidak_cocok']),
                'contoh': v_pdf2json['tidak_cocok'][:5],
            })

        # Peringatan format sumber (bukan kesalahan ekstraksi)
        for w in unit['validasi'].get('peringatan_format_sumber', []):
            peringatan_sumber.append({
                'kode_unit': job['kode'],
                'kategori': w.get('kategori'),
                'nomor_kuk': w.get('nomor_kuk'),
                'halaman_pdf': w.get('halaman_pdf'),
                'pesan': w.get('pesan'),
            })

        if info['level_melebihi_tiga']:
            masalah_format.append({
                'kode_unit': job['kode'],
                'masalah': f"penomoran level {info['levels']} melebihi 3 tingkat",
                'tindakan': 'level dipertahankan, perlu penentuan format Word',
            })

        baris_laporan.append({
            'SKKNI': job['skkni'],
            'kode_unit': job['kode'],
            'elemen': len(unit['elemen_kompetensi']),
            'kuk': sum(len(e['kuk']) for e in unit['elemen_kompetensi']),
            'validasi': unit['validasi']['status'],
            'docx': str(docx_path.relative_to(project_root)),
            'pdf_render': str(pdf_out.relative_to(project_root)) if dirender else None,
            'halaman_pdf_sumber': target['end_page'] - target['start_page'] + 1,
            'halaman_render': v_docx2pdf['halaman'],
            'paragraf_bernomor': v_json2docx.get('jumlah_paragraf_bernomor'),
            'paragraf_tanpa_nomor': v_json2docx.get('jumlah_paragraf_tanpa_nomor'),
            'tabel': v_json2docx.get('jumlah_tabel'),
            'perlu_verifikasi': v_json['perlu_verifikasi'],
            'catatan': v_json['catatan'],
        })
        hasil_unit.append({'job': job, 'unit': unit, 'info': info})

    # ------------------------------------------------------------- regresi
    print(f"\n{'=' * 100}\nREGRESI\n{'=' * 100}")
    regresi = {}
    for nama in ('test_extractor.py', 'test_unit_parser.py', 'test_unit_complete.py',
                 'test_unit_numbering.py', 'test_docx_generator.py'):
        r = subprocess.run([sys.executable, str(project_root / 'tests' / nama)],
                           capture_output=True, text=True, cwd=str(project_root))
        ringkas = next((l for l in reversed(r.stdout.splitlines()) if 'HASIL:' in l), '')
        regresi[nama] = {'exit': r.returncode, 'ringkas': ringkas.strip()}
        cek('regresi', f'{nama} lulus', r.returncode == 0, ringkas)

    # -------------------------------------------------------------- laporan
    laporan = {
        'dibuat': datetime.now().isoformat(timespec='seconds'),
        'tabel': baris_laporan,
        'ringkasan': {
            'unit_diproses': len(baris_laporan),
            'pemeriksaan_berhasil': len(lulus),
            'pemeriksaan_gagal': len(gagal),
            'bagian_perlu_verifikasi': sum(b['perlu_verifikasi'] for b in baris_laporan),
            'peringatan_format_sumber': len(peringatan_sumber),
        },
        'gagal': [{'unit': u, 'pemeriksaan': n} for u, n in gagal],
        'perbedaan_teks': perbedaan_teks,
        'peringatan_format_sumber': peringatan_sumber,
        'masalah_format_word': masalah_format,
        'perbaikan_parser': [
            'Nomor halaman tidak lagi terbaca sebagai nomor bagian (angka telanjang ditolak).',
            'Deteksi nomor halaman diperlebar (sampai 110 pt dari tepi bawah).',
            'Batas tabel ditentukan dari posisi header & BATASAN VARIABEL, bukan urutan blok.',
            'Blok yang menyeberangi dua kolom dipecah per baris (KUK tidak lagi hilang).',
            'Nomor yang berdiri sendiri di satu baris dipasangkan dengan teks baris berikutnya.',
            'Kode unit lain di dalam bagian tidak lagi dianggap kebocoran (hanya label KODE UNIT).',
            'Generator DOCX mendukung level penomoran >3 tanpa memotongnya.',
        ],
        'regresi': regresi,
    }
    REPORT.write_text(json.dumps(laporan, ensure_ascii=False, indent=2), encoding='utf-8')

    # --------------------------------------------------------------- tabel
    print("\n" + "=" * 100)
    print(f"{'SKKNI':<16} {'Kode Unit':<20} {'El':>3} {'KUK':>4} {'Validasi':<18} {'DOCX':<10}")
    print("-" * 100)
    for b in baris_laporan:
        print(f"{b['SKKNI']:<16} {b['kode_unit']:<20} {b['elemen']:>3} {b['kuk']:>4} "
              f"{b['validasi']:<18} {'ada' if b['docx'] else '-':<10}")
    print("=" * 100)
    print(f"HASIL: {len(lulus)} lulus, {len(gagal)} gagal")
    if gagal:
        for u, n in gagal:
            print(f"  - [{u}] {n}")
    print(f"Laporan: {REPORT.relative_to(project_root)}")
    print("=" * 100)
    return 0 if not gagal else 1


if __name__ == '__main__':
    sys.exit(main())
