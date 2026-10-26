"""
Pengujian generator DOCX (Fase 3A).

Memeriksa isi dan struktur XML, bukan hanya teks hasil ekstraksi:
- identitas, tabel, jumlah elemen/KUK
- numbering OOXML (abstractNum, num, startOverride, ilvl, ind)
- paragraf tanpa nomor
- kesesuaian teks dengan JSON (tidak hilang/duplikat)
- render PDF untuk pemeriksaan tampilan
"""

import json
import re
import shutil
import subprocess
import sys
import zipfile
from collections import Counter
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docx.oxml.ns import qn
import pymupdf as fitz

from app.docx_generator import (
    generate_docx, LEVEL_INDENT_TWIPS, HANGING_TWIPS, FONT_NAME, FONT_SIZE_PT,
    PAGE_W_TWIPS, PAGE_H_TWIPS, MARGIN_TOP, MARGIN_RIGHT, MARGIN_BOTTOM,
    MARGIN_LEFT, COL_WIDTH_TWIPS,
)
from app.section_parser import flatten_nodes, normalize_text

JSON_SRC = project_root / 'data' / 'json' / 'test_unit_001_complete.json'
OUT_DIR = project_root / 'output'
DOCX_PATH = OUT_DIR / 'M.73ADV00.001.1.docx'
PDF_PATH = OUT_DIR / 'M.73ADV00.001.1.pdf'

lulus, gagal = [], []


def cek(nama, kondisi, detail=''):
    (lulus if kondisi else gagal).append(nama)
    print(f"  [{'PASS' if kondisi else 'FAIL'}] {nama}" + (f" — {detail}" if detail and not kondisi else ''))


def teks_rapat(s: str) -> str:
    return re.sub(r'\s+', '', s)


def main():
    print("=" * 100)
    print("PENGUJIAN GENERATOR DOCX — FASE 3A")
    print("=" * 100)

    unit = json.loads(JSON_SRC.read_text(encoding='utf-8'))
    info = generate_docx(str(JSON_SRC), str(DOCX_PATH))
    print(f"\nDOCX: {DOCX_PATH.relative_to(project_root)} "
          f"({DOCX_PATH.stat().st_size / 1024:.0f} KB)")

    bv_nodes = flatten_nodes(unit['batasan_variabel']['sections'])
    pp_nodes = flatten_nodes(unit['panduan_penilaian']['sections'])
    bv_bernomor = [n for n in bv_nodes if n['is_numbered']]
    pp_bernomor = [n for n in pp_nodes if n['is_numbered']]
    semua_bernomor = bv_bernomor + pp_bernomor
    semua_tanpa = [n for n in bv_nodes + pp_nodes if not n['is_numbered']]
    jumlah_kuk = sum(len(e['kuk']) for e in unit['elemen_kompetensi'])

    # ------------------------------------------------------------ isi dasar
    print("\n1. ISI DOKUMEN (python-docx)")
    doc = Document(str(DOCX_PATH))
    teks_dok = '\n'.join(p.text for p in doc.paragraphs)
    for tabel in doc.tables:
        for row in tabel.rows:
            for cell in row.cells:
                teks_dok += '\n' + '\n'.join(p.text for p in cell.paragraphs)

    cek('KODE UNIT tercetak', f"KODE UNIT: {unit['kode_unit']}" in teks_dok)
    cek('JUDUL UNIT tercetak', unit['judul_unit'] in teks_dok)
    cek('DESKRIPSI UNIT tercetak', teks_rapat(unit['deskripsi_unit']) in teks_rapat(teks_dok))
    cek('heading BATASAN VARIABEL ada',
        any(p.text.strip() == 'BATASAN VARIABEL' for p in doc.paragraphs))
    cek('heading PANDUAN PENILAIAN ada', any(p.text.strip() == 'PANDUAN PENILAIAN'
                                              for p in doc.paragraphs))

    # Regresi audit Fase 3B: deskripsi tidak boleh menyerap bagian lain
    label_struktur = ('BATASAN VARIABEL', 'PANDUAN PENILAIAN', 'ELEMEN KOMPETENSI',
                      'KRITERIA UNJUK KERJA', 'KODE UNIT', 'JUDUL UNIT')
    cek('deskripsi bebas label struktur',
        not any(l in unit['deskripsi_unit'].upper() for l in label_struktur))
    cek('deskripsi tercetak utuh satu paragraf',
        sum(1 for p in doc.paragraphs
            if normalize_text(p.text) == normalize_text(unit['deskripsi_unit'])) == 1)
    cek('heading bagian muncul sekali',
        sum(1 for p in doc.paragraphs if p.text.strip() == 'BATASAN VARIABEL') == 1
        and sum(1 for p in doc.paragraphs if p.text.strip() == 'PANDUAN PENILAIAN') == 1)

    # --------------------------------------------------------------- tabel
    print("\n2. TABEL ELEMEN / KUK")
    cek('tepat 1 tabel Word', len(doc.tables) == 1, str(len(doc.tables)))
    tabel = doc.tables[0]
    cek('2 kolom', len(tabel.columns) == 2, str(len(tabel.columns)))
    cek('jumlah baris = 1 header + 4 elemen', len(tabel.rows) == 5, str(len(tabel.rows)))

    hdr = tabel.rows[0]
    cek('header kiri "ELEMEN KOMPETENSI"',
        hdr.cells[0].text.strip() == 'ELEMEN KOMPETENSI', hdr.cells[0].text)
    cek('header kanan "KRITERIA UNJUK KERJA"',
        hdr.cells[1].text.strip() == 'KRITERIA UNJUK KERJA', hdr.cells[1].text)
    bold_hdr = all(r.bold for c in hdr.cells for p in c.paragraphs for r in p.runs)
    cek('header tabel dicetak tebal', bold_hdr)

    kuk_terkumpul = []
    for i, e in enumerate(unit['elemen_kompetensi'], start=1):
        baris = tabel.rows[i]
        kiri = normalize_text(baris.cells[0].text)
        cek(f'elemen {e["nomor_elemen"]} di baris sendiri', kiri.startswith(e['nomor_elemen'])
            and teks_rapat(e['nama_elemen']) in teks_rapat(kiri), kiri[:40])
        kanan = baris.cells[1]
        nomor_kuk = [p.text.split('\t', 1)[0].strip() for p in kanan.paragraphs]
        cek(f'KUK elemen {e["nomor_elemen"]} lengkap dan berurutan',
            nomor_kuk == [k['nomor_kuk'] for k in e['kuk']], str(nomor_kuk))
        kuk_terkumpul.extend(kanan.paragraphs)

    cek('total KUK sesuai JSON', len(kuk_terkumpul) == jumlah_kuk,
        f'{len(kuk_terkumpul)} vs {jumlah_kuk}')
    for p, k in zip(kuk_terkumpul, [k for e in unit['elemen_kompetensi'] for k in e['kuk']]):
        if not p.text.split('\t', 1)[-1].strip() == k['uraian_kuk']:
            cek(f'uraian KUK {k["nomor_kuk"]} identik', False, p.text[:60])
            break
    else:
        cek('seluruh uraian KUK identik dengan JSON', True)

    # ---------------------------------------------------------- numbering
    print("\n3. STRUKTUR NUMBERING (OOXML)")
    with zipfile.ZipFile(DOCX_PATH) as z:
        numbering_xml = z.read('word/numbering.xml').decode('utf-8')
        document_xml = z.read('word/document.xml').decode('utf-8')

    num_ids = info['num_ids']
    abs_id = info['abstract_num_id']
    cek('abstractNum baru terdaftar', f'w:abstractNumId="{abs_id}"' in numbering_xml)
    for lvl, teks in ((0, '%1.'), (1, '%1.%2.'), (2, '%1.%2.%3.')):
        pola = rf'<w:lvl w:ilvl="{lvl}"[^>]*>.*?<w:lvlText w:val="{re.escape(teks)}"'
        cek(f'level {lvl + 1} memakai format {teks}',
            re.search(pola, numbering_xml, re.S) is not None)

    for grup, nid in num_ids.items():
        blok = re.search(rf'<w:num w:numId="{nid}".*?</w:num>', numbering_xml, re.S)
        cek(f'numId {nid} ({grup}) ada', blok is not None)
        if blok:
            cek(f'numId {nid} punya 3 lvlOverride startOverride=1',
                blok.group(0).count('<w:lvlOverride ') == 3
                and blok.group(0).count('<w:startOverride w:val="1"/>') == 3,
                f"override={blok.group(0).count('<w:lvlOverride ')} "
                f"start={blok.group(0).count('<w:startOverride w:val="1"/>')}")

    cek('dua kelompok memakai numId berbeda', num_ids['batasan_variabel'] != num_ids['panduan_penilaian'])

    # ------------------------------------------------- paragraf bernomor
    print("\n4. PARAGRAF BERNOMOR")
    ber_nomor = []
    for p in doc.paragraphs:
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn('w:numPr')) is not None:
            numPr = pPr.find(qn('w:numPr'))
            ilvl = int(numPr.find(qn('w:ilvl')).get(qn('w:val')))
            nid = int(numPr.find(qn('w:numId')).get(qn('w:val')))
            ind = pPr.find(qn('w:ind'))
            ber_nomor.append({
                'p': p, 'ilvl': ilvl, 'numId': nid,
                'left': int(ind.get(qn('w:left'))) if ind is not None else None,
                'hanging': int(ind.get(qn('w:hanging'))) if ind is not None
                           and ind.get(qn('w:hanging')) else 0,
                'teks': p.text,
            })

    cek('jumlah paragraf bernomor = 51', len(ber_nomor) == 51, str(len(ber_nomor)))
    cek('semua numId termasuk dua kelompok', {b['numId'] for b in ber_nomor} == set(num_ids.values()))
    cek('semua ilvl 0..2', all(0 <= b['ilvl'] <= 2 for b in ber_nomor))

    # ------------------------------------------------------- paragraf biasa
    print("\n5. PARAGRAF TANPA NOMOR")
    tanpa_nomor = [p for p in doc.paragraphs
                   if (p._p.pPr is None or p._p.pPr.find(qn('w:numPr')) is None)
                   and p.text.strip() in {n['teks'] for n in semua_tanpa}]
    cek('3 paragraf "(Tidak ada.)" tanpa numPr', len(tanpa_nomor) == 3, str(len(tanpa_nomor)))

    # Pasangkan urutan paragraf di DOCX dengan urutan node JSON
    for p, node in zip(tanpa_nomor, semua_tanpa):
        ind = p._p.pPr.find(qn('w:ind')) if p._p.pPr is not None else None
        left = int(ind.get(qn('w:left'))) if ind is not None and ind.get(qn('w:left')) else 0
        hanging = ind.get(qn('w:hanging')) if ind is not None else None
        induk_level = len(node['parent_nomor'].split('.'))
        harap = LEVEL_INDENT_TWIPS.get(induk_level + 1, LEVEL_INDENT_TWIPS[3])
        cek(f'paragraf tanpa nomor (induk {node["parent_nomor"]}) indent {harap} tanpa hanging',
            p.text.strip() == node['teks'] and left == harap and hanging is None,
            f'left={left} harap={harap} hanging={hanging} teks={p.text!r}')

    # ------------------------------------------------------ indentasi level
    print("\n6. INDENTASI SESUAI LEVEL")
    salah_ind = []
    for b in ber_nomor:
        harap = LEVEL_INDENT_TWIPS[b['ilvl'] + 1]
        harap_hang = info['hanging_per_level'].get(b['ilvl'] + 1, HANGING_TWIPS)
        if b['left'] != harap or b['hanging'] != harap_hang:
            salah_ind.append((b['teks'][:25], b['ilvl'], b['left'], b['hanging']))
    cek('left & hanging setiap level sesuai tabel target', not salah_ind, str(salah_ind[:3]))
    cek('tidak memakai atribut leftChars/firstLineChars (membuat indent diabaikan)',
        'w:leftChars' not in document_xml and 'w:firstLineChars' not in document_xml)

    # ------------------------------------------------------- urutan nomor
    print("\n7. NOMOR YANG DITAMPILKAN SESUAI JSON")
    counter = {nid: [0, 0, 0] for nid in num_ids.values()}
    tampil = []
    for b in ber_nomor:
        ilvl, nid = b['ilvl'], b['numId']
        counter[nid][ilvl] += 1
        for lebih_dalam in range(ilvl + 1, 3):
            counter[nid][lebih_dalam] = 0
        tampil.append('.'.join(str(counter[nid][i]) for i in range(ilvl + 1)))

    harap_nomor = [n['nomor'] for n in semua_bernomor]
    cek('urutan tampilan numbering = nomor JSON', tampil == harap_nomor,
        str([(a, b) for a, b in zip(tampil, harap_nomor) if a != b][:5]))

    awal_bv = ber_nomor[0]
    awal_pp = next(b for b in ber_nomor if b['numId'] == num_ids['panduan_penilaian'])
    cek('BV dimulai dari 1 (ilvl 0)', awal_bv['ilvl'] == 0 and awal_bv['teks'] == semua_bernomor[0]['teks'])
    cek('PP dimulai ulang dari 1 (ilvl 0)',
        awal_pp['ilvl'] == 0 and tampil[ber_nomor.index(awal_pp)] == '1',
        tampil[ber_nomor.index(awal_pp)])

    # ---------------------------------------------------------- tabel XML
    print("\n8. STRUKTUR TABEL XML")
    cek('baris header diulang antar halaman (tblHeader)', '<w:tblHeader/>' in document_xml)
    cek('semua baris cantSplit', document_xml.count('<w:cantSplit/>') == 5,
        str(document_xml.count('<w:cantSplit/>')))
    grids = re.findall(r'<w:gridCol w:w="(\d+)"', document_xml)
    cek('lebar kolom 4524 twip', grids == [str(COL_WIDTH_TWIPS)] * 2, str(grids))

    # ------------------------------------------------------------- teks
    print("\n9. KELENGKAPAN TEKS (tidak hilang / ganda)")
    teks_dok_rapat = teks_rapat(teks_dok)
    hilang = []
    for teks in [unit['judul_unit'], unit['deskripsi_unit']] + \
                [n['teks'] for n in semua_bernomor + semua_tanpa] + \
                [e['nama_elemen'] for e in unit['elemen_kompetensi']] + \
                [k['uraian_kuk'] for e in unit['elemen_kompetensi'] for k in e['kuk']]:
        if teks_rapat(teks) not in teks_dok_rapat:
            hilang.append(teks[:50])
    cek('seluruh teks JSON muncul di DOCX', not hilang, str(hilang[:3]))

    nomor_json = [k['nomor_kuk'] for e in unit['elemen_kompetensi'] for k in e['kuk']]
    nomor_docx = [p.text.split('\t', 1)[0].strip() for p in kuk_terkumpul]
    cek('nomor KUK tidak berubah/digandakan', nomor_docx == nomor_json,
        str([(a, b) for a, b in zip(nomor_docx, nomor_json) if a != b][:3]))

    # --------------------------------------------------------- tipografi
    print("\n10. TIPOGRAFI & HALAMAN")
    normal = doc.styles['Normal']
    cek('font Bookman Old Style 11pt',
        normal.font.name == FONT_NAME and normal.font.size.pt == FONT_SIZE_PT,
        f'{normal.font.name} {normal.font.size}')
    sec = doc.sections[0]
    cek('ukuran halaman F4 sesuai referensi',
        sec.page_width.twips == PAGE_W_TWIPS and sec.page_height.twips == PAGE_H_TWIPS)
    cek('margin sesuai referensi',
        (sec.top_margin.twips, sec.right_margin.twips, sec.bottom_margin.twips,
         sec.left_margin.twips) == (MARGIN_TOP, MARGIN_RIGHT, MARGIN_BOTTOM, MARGIN_LEFT),
        str((sec.top_margin.twips, sec.right_margin.twips,
             sec.bottom_margin.twips, sec.left_margin.twips)))

    # ------------------------------------------------------------ render
    print("\n11. RENDER PDF (LibreOffice)")
    pdf_info = {'tersedia': False}
    if shutil.which('flatpak'):
        r = subprocess.run(
            ['flatpak', 'run', 'org.libreoffice.LibreOffice', '--headless',
             '--convert-to', 'pdf', '--outdir', str(OUT_DIR), str(DOCX_PATH)],
            capture_output=True, text=True, timeout=240)
        pdf_info['tersedia'] = PDF_PATH.exists()
        pdf_info['stdout'] = (r.stdout or '').strip().splitlines()[-1:] or []
    if pdf_info['tersedia']:
        pdf = fitz.open(str(PDF_PATH))
        pdf_info['halaman'] = pdf.page_count
        teks_pdf = '\n'.join(p.get_text() for p in pdf)
        pdf_rapat = teks_rapat(teks_pdf)
        cek('PDF hasil render terbentuk', pdf.page_count > 0, str(pdf.page_count))
        cek('penomoran tampil di PDF (2.1.1, 3.1.10, 4.2.1)',
            all(x in teks_pdf for x in ('2.1.1', '3.1.10', '4.2.1')))
        hilang_pdf = [n['nomor'] for n in semua_bernomor
                      if teks_rapat(n['teks'])[:40] not in pdf_rapat]
        cek('seluruh teks bagian tampil di PDF', not hilang_pdf, str(hilang_pdf[:5]))
        cek('"(Tidak ada.)" tampil 3 kali tanpa nomor',
            len(re.findall(r'(?m)^\(Tidak ada\.\)', teks_pdf)) == 3,
            str(len(re.findall(r'(?m)^\(Tidak ada\.\)', teks_pdf))))

        # Posisi nyata nomor di PDF: level 1 di margin, level 2 +21 pt, level 3 +42 pt
        posisi = {}
        for page in pdf:
            for w in page.get_text('words'):
                t = w[4]
                if t == '1.1.' and 'pos_l2' not in posisi:
                    posisi['pos_l2'] = w[0]
                if t == '2.1.1.' and 'pos_l3' not in posisi:
                    posisi['pos_l3'] = w[0]
                if t == '1.' and 'pos_l1' not in posisi:
                    posisi['pos_l1'] = w[0]

        # Nomor dimulai di (left indent - hanging); hanging kini adaptif
        exp_l2 = (LEVEL_INDENT_TWIPS[2] - info['hanging_per_level'][2]) / 20.0
        exp_l3 = (LEVEL_INDENT_TWIPS[3] - info['hanging_per_level'][3]) / 20.0
        cek(f'posisi nomor level 2 sesuai hanging ({exp_l2:.1f} pt dari level 1)',
            posisi.get('pos_l2') is not None and posisi.get('pos_l1') is not None
            and abs((posisi['pos_l2'] - posisi['pos_l1']) - exp_l2) < 2,
            str(posisi))
        cek(f'posisi nomor level 3 sesuai hanging ({exp_l3:.1f} pt dari level 1)',
            posisi.get('pos_l3') is not None and posisi.get('pos_l1') is not None
            and abs((posisi['pos_l3'] - posisi['pos_l1']) - exp_l3) < 2,
            str(posisi))

        # Jarak nomor ke teks (bukan hanya properti XML)
        jarak = []
        for page in pdf:
            for b in page.get_text('dict')['blocks']:
                for l in b.get('lines', []):
                    sp = l['spans']
                    if len(sp) >= 3 and re.fullmatch(r'\d+(?:\.\d+)+\.', sp[0]['text'].strip()):
                        jarak.append(sp[-1]['bbox'][0] - sp[0]['bbox'][2])
        cek('jarak nomor-teks 4-12 pt (tidak menempel / melompat)',
            bool(jarak) and all(4 <= g <= 12 for g in jarak),
            str([round(g, 1) for g in jarak[:6]]))
        cek('jarak nomor-teks konsisten (selisih <= 2 pt)',
            bool(jarak) and max(jarak) - min(jarak) <= 2,
            str([round(g, 1) for g in jarak[:6]]))
        pdf.close()
    else:
        print("  [SKIP] LibreOffice tidak tersedia untuk render PDF")

    # ---------------------------------------------------------- ringkasan
    print("\n" + "-" * 100)
    print("RINGKASAN")
    print(f"  Elemen / KUK        : {len(unit['elemen_kompetensi'])} / {jumlah_kuk}")
    print(f"  Paragraf bernomor   : {len(ber_nomor)} (BV {len(bv_bernomor)} + PP {len(pp_bernomor)})")
    print(f"  Paragraf tanpa nomor: {len(tanpa_nomor)}")
    print(f"  Tabel               : {len(doc.tables)}")
    print(f"  numId               : {num_ids}")
    print(f"  Render PDF          : {pdf_info.get('halaman', 'n/a')} halaman")

    print("\n" + "=" * 100)
    print(f"HASIL: {len(lulus)} lulus, {len(gagal)} gagal")
    if gagal:
        print("GAGAL:")
        for g in gagal:
            print(f"  - {g}")
    print("=" * 100)
    return 0 if not gagal else 1


if __name__ == '__main__':
    sys.exit(main())
