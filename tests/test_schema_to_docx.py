"""
Pengujian Skema -> DOCX (Fase 3D).

Memeriksa dokumen gabungan 10 unit: urutan, kelengkapan, tabel, numbering,
pemisahan halaman, jarak nomor-teks, dan kesesuaian dengan JSON + PDF sumber.
"""

import json
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from docx.oxml.ns import qn
import pymupdf as fitz

from app.section_parser import flatten_nodes, normalize_text
from app.schema_pipeline import proses, NAMA_FINAL

JSON_DIR = project_root / 'data' / 'json' / 'schema_content_creator'
SKEMA_JSON = project_root / 'data' / 'json' / 'schema_content_creator.json'
REPORT = project_root / 'output' / 'laporan_content_creator.json'
DOCX = project_root / 'output' / NAMA_FINAL
PDF = DOCX.with_suffix('.pdf')

lulus, gagal = [], []


def cek(nama, kondisi, detail=''):
    (lulus if kondisi else gagal).append(nama)
    print(f"  [{'PASS' if kondisi else 'FAIL'}] {nama}" + (f" — {detail}" if detail and not kondisi else ''))


def rapat(s):
    return re.sub(r'\s+', '', s or '')


def main():
    print("=" * 100)
    print("PENGUJIAN SKEMA -> DOCX — FASE 3D")
    print("=" * 100)

    # Pastikan artefak terbaru (pipeline idempoten)
    laporan = proses()
    print(f"\nDOCX: {DOCX.relative_to(project_root)} | halaman: {laporan['ringkasan']['halaman_docx']}")

    skema = json.loads(SKEMA_JSON.read_text(encoding='utf-8'))
    units = [json.loads((JSON_DIR / f"{u['kode_unit']}.json").read_text(encoding='utf-8'))
             for u in skema['units']]

    doc = Document(str(DOCX))
    paras = doc.paragraphs
    teks_doc = '\n'.join(p.text for p in paras)
    teks_lengkap = teks_doc
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                teks_lengkap += '\n' + '\n'.join(p.text for p in c.paragraphs)

    # ------------------------------------------------------------ urutan unit
    print("\n1. URUTAN & IDENTITAS")
    kode_paras = [p.text.strip() for p in paras if p.text.strip().startswith('KODE UNIT:')]
    cek('10 unit dalam dokumen', len(kode_paras) == 10, str(len(kode_paras)))
    urutan = [k.replace('KODE UNIT:', '').strip() for k in kode_paras]
    cek('urutan sesuai dokumen skema', urutan == [u['kode_unit'] for u in skema['units']],
        str(urutan))
    for u in units:
        cek(f"{u['kode_unit']}: identitas lengkap",
            u['judul_unit'] in teks_doc and rapat(u['deskripsi_unit'])[:60] in rapat(teks_lengkap))

    # ---------------------------------------------------------------- tabel
    print("\n2. TABEL & KUK")
    cek('10 tabel Word', len(doc.tables) == 10, str(len(doc.tables)))
    total_kuk = 0
    for tabel, u in zip(doc.tables, units):
        kuk_json = [k['nomor_kuk'] for e in u['elemen_kompetensi'] for k in e['kuk']]
        kuk_docx = []
        for r in tabel.rows[1:]:
            kuk_docx.extend(p.text.split('\t', 1)[0].strip() for p in r.cells[1].paragraphs)
        cek(f"{u['kode_unit']}: tabel 2 kolom & baris sesuai",
            len(tabel.columns) == 2 and len(tabel.rows) == 1 + len(u['elemen_kompetensi']),
            f"{len(tabel.columns)}x{len(tabel.rows)}")
        cek(f"{u['kode_unit']}: nomor KUK lengkap", kuk_docx == kuk_json,
            str([(a, b) for a, b in zip(kuk_docx, kuk_json) if a != b][:3]))
        total_kuk += len(kuk_json)
    cek('total KUK sesuai laporan', total_kuk == laporan['ringkasan']['total_kuk'],
        f"{total_kuk} vs {laporan['ringkasan']['total_kuk']}")

    # ------------------------------------------------------------ teks bagian
    print("\n3. KELENGKAPAN BATASAN VARIABEL & PANDUAN PENILAIAN")
    hilang_total = []
    for u in units:
        nodes = (flatten_nodes(u['batasan_variabel']['sections'])
                 + flatten_nodes(u['panduan_penilaian']['sections']))
        hilang = [n['nomor'] for n in nodes if rapat(n['teks'])[:40] not in rapat(teks_lengkap)]
        if hilang:
            hilang_total.append((u['kode_unit'], hilang[:5]))
    cek('seluruh teks 10 unit ada', not hilang_total, str(hilang_total[:3]))
    cek('heading muncul sekali per unit',
        teks_doc.count('BATASAN VARIABEL') == 10 and teks_doc.count('PANDUAN PENILAIAN') == 10,
        f"{teks_doc.count('BATASAN VARIABEL')}/{teks_doc.count('PANDUAN PENILAIAN')}")

    # ------------------------------------------------------------- numbering
    print("\n4. NUMBERING & RESTART")
    ber_nomor = []
    for p in paras:
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn('w:numPr')) is not None:
            numPr = pPr.find(qn('w:numPr'))
            ber_nomor.append((int(numPr.find(qn('w:numId')).get(qn('w:val'))),
                              int(numPr.find(qn('w:ilvl')).get(qn('w:val')))))
    num_ids = {nid for nid, _ in ber_nomor}
    cek('20 numId unik (2 per unit)', len(num_ids) == 20, str(len(num_ids)))

    # simulasi counter per numId
    counter = {nid: [0, 0, 0, 0] for nid in num_ids}
    tampil = []
    for nid, ilvl in ber_nomor:
        counter[nid][ilvl] += 1
        for d in range(ilvl + 1, 4):
            counter[nid][d] = 0
        tampil.append('.'.join(str(counter[nid][i]) for i in range(ilvl + 1)))

    idx = 0
    ok_nomor = True
    ok_restart = []
    for u in units:
        nodes = [n for n in (flatten_nodes(u['batasan_variabel']['sections'])
                             + flatten_nodes(u['panduan_penilaian']['sections']))
                 if n['is_numbered']]
        n_bv = len([n for n in flatten_nodes(u['batasan_variabel']['sections'])
                    if n['is_numbered']])
        bagian = tampil[idx:idx + len(nodes)]
        idx += len(nodes)
        if bagian != [n['nomor'] for n in nodes]:
            ok_nomor = False
        ok_restart.append(len(bagian) > n_bv and bagian[0] == '1' and bagian[n_bv] == '1')
    cek('nomor tampilan = nomor JSON untuk 10 unit', ok_nomor)
    cek('BV & PP tiap unit mulai dari 1', all(ok_restart),
        str([units[i]['kode_unit'] for i, v in enumerate(ok_restart) if not v]))

    # -------------------------------------------------------- halaman & render
    print("\n5. RENDER PDF")
    cek('PDF tersedia', PDF.exists())
    if PDF.exists():
        pdf = fitz.open(str(PDF))
        cek('halaman > 0', pdf.page_count > 0, str(pdf.page_count))

        halaman_unit = {}
        for i, page in enumerate(pdf):
            m = re.search(r'KODE UNIT: (\S+)', page.get_text())
            if m:
                halaman_unit[m.group(1)] = i + 1
                y = next((l['bbox'][1] for b in page.get_text('dict')['blocks']
                          for l in b.get('lines', [])
                          if ''.join(s['text'] for s in l['spans']).strip().startswith('KODE UNIT:')), 999)
                if y >= 120:
                    cek(f"{m.group(1)}: mulai di atas halaman", False, f"y={y}")
        cek('10 unit punya halaman awal', len(halaman_unit) == 10, str(len(halaman_unit)))
        cek('urutan halaman menaik',
            [halaman_unit.get(u['kode_unit']) for u in units]
            == sorted(halaman_unit.get(u['kode_unit']) for u in units))
        kosong = [i + 1 for i, p in enumerate(pdf) if len(p.get_text().strip()) < 40]
        cek('tidak ada halaman kosong', not kosong, str(kosong))

        teks_pdf = rapat('\n'.join(p.get_text() for p in pdf))
        hilang_pdf = []
        for u in units:
            for n in (flatten_nodes(u['batasan_variabel']['sections'])
                      + flatten_nodes(u['panduan_penilaian']['sections'])):
                if rapat(n['teks'])[:40] not in teks_pdf:
                    hilang_pdf.append((u['kode_unit'], n['nomor']))
        cek('seluruh teks tampil di PDF', not hilang_pdf, str(hilang_pdf[:3]))

        jarak = []
        for page in pdf:
            for b in page.get_text('dict')['blocks']:
                for l in b.get('lines', []):
                    sp = l['spans']
                    if len(sp) >= 3 and re.fullmatch(r'\d+(?:\.\d+)+\.', sp[0]['text'].strip()):
                        jarak.append(sp[-1]['bbox'][0] - sp[0]['bbox'][2])
        cek('jarak nomor-teks 4-12 pt', bool(jarak) and all(4 <= g <= 12 for g in jarak),
            f"min={min(jarak):.1f} max={max(jarak):.1f}" if jarak else 'n/a')
        pdf.close()

    # ---------------------------------------------------------------- laporan
    print("\n6. LAPORAN")
    lap = json.loads(REPORT.read_text(encoding='utf-8'))
    cek('laporan tersimpan', lap['ringkasan']['unit_skema'] == 10)
    cek('dokumen dinyatakan lengkap', lap['ringkasan']['dokumen_lengkap'] is True)
    cek('nama dokumen final (bukan DRAFT)', 'DRAFT' not in lap['ringkasan']['nama_docx'])
    cek('10 unit diekstrak', lap['ringkasan']['unit_diekstrak'] == 10)
    cek('semua unit valid_struktur', lap['ringkasan']['unit_perlu_verifikasi'] == 0,
        str([h['kode_unit'] for h in lap['tabel'] if h['validasi'] != 'valid_struktur']))
    cek('tidak ada kegagalan ekstraksi', not lap['gagal_ekstrak'])

    # ---------------------------------------------------------------- regresi
    print("\n7. REGRESI FASE SEBELUMNYA")
    for nama in ('test_extractor.py', 'test_unit_parser.py', 'test_unit_complete.py',
                 'test_unit_numbering.py', 'test_docx_generator.py',
                 'test_cross_skkni.py', 'test_multi_docx.py'):
        r = subprocess.run([sys.executable, str(project_root / 'tests' / nama)],
                           capture_output=True, text=True, cwd=str(project_root))
        cek(f'{nama} lulus', r.returncode == 0,
            (r.stdout + r.stderr).strip().splitlines()[-1] if r.returncode else '')

    print("\n" + "-" * 100)
    print(f"Unit: {len(units)} | Elemen: {laporan['ringkasan']['total_elemen']} | "
          f"KUK: {laporan['ringkasan']['total_kuk']} | "
          f"Halaman: {laporan['ringkasan']['halaman_docx']}")
    print("\n" + "=" * 100)
    print(f"HASIL: {len(lulus)} lulus, {len(gagal)} gagal")
    if gagal:
        for g in gagal:
            print(f"  - {g}")
    print("=" * 100)
    return 0 if not gagal else 1




def test_schema_to_docx():
    assert main() == 0


if __name__ == '__main__':
    sys.exit(main())
