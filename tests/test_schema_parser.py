"""
Pengujian Schema Parser (Fase 3D).

Membaca dokumen skema "21. Content Creator.docx" dan memastikan daftar unit
diambil dari tabel bagian Kemasan/Paket Kompetensi, bukan dari bagian lain.
"""

import json
import re
import shutil
import sys
from pathlib import Path

project_root = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(project_root))

from docx import Document
from app.schema_parser import baca_skema, simpan_skema, UNIT_CODE_PATTERN

SKEMA = project_root / 'data' / 'skema' / '21. Content Creator.docx'
OUTPUT = project_root / 'data' / 'json' / 'schema_content_creator.json'

lulus, gagal = [], []


def cek(nama, kondisi, detail=''):
    (lulus if kondisi else gagal).append(nama)
    print(f"  [{'PASS' if kondisi else 'FAIL'}] {nama}" + (f" — {detail}" if detail and not kondisi else ''))


def main():
    print("=" * 100)
    print("PENGUJIAN SCHEMA PARSER — FASE 3D")
    print("=" * 100)

    if not SKEMA.exists():
        print(f"ERROR: dokumen skema tidak ditemukan: {SKEMA}")
        return 1

    hasil = baca_skema(SKEMA)
    simpan_skema(hasil, OUTPUT)

    print(f"\nDokumen: {hasil['sumber']} | unit: {hasil['jumlah_unit']} "
          f"| tabel: {hasil['tabel_index']}")

    # ---------------------------------------------------------------- jumlah
    print("\n1. JUMLAH & URUTAN")
    cek('10 unit terbaca', hasil['jumlah_unit'] == 10, str(hasil['jumlah_unit']))
    cek('nomor urut 1..10',
        [u['nomor_urut'] for u in hasil['units']] == list(range(1, 11)),
        str([u['nomor_urut'] for u in hasil['units']]))

    # ------------------------------------------------------------- validitas
    print("\n2. KODE & SUMBER")
    kode = [u['kode_unit'] for u in hasil['units']]
    cek('semua kode unik', len(kode) == len(set(kode)))
    cek('semua kode sesuai pola SKKNI',
        all(UNIT_CODE_PATTERN.fullmatch(k) for k in kode),
        str([k for k in kode if not UNIT_CODE_PATTERN.fullmatch(k)]))
    cek('semua unit mencatat sumber tabel',
        all(u['sumber']['tabel_index'] == hasil['tabel_index'] for u in hasil['units']))
    cek('baris tabel menaik', [u['sumber']['baris_tabel'] for u in hasil['units']]
        == sorted(u['sumber']['baris_tabel'] for u in hasil['units']))

    # Pastikan kode tidak diambil dari luar tabel (mis. daftar acuan normatif)
    sumber_doc = Document(str(SKEMA))
    tabel = sumber_doc.tables[hasil['tabel_index']]
    kode_tabel = [re.sub(r'\s+', ' ', r.cells[1].text).strip() for r in tabel.rows[1:]]
    cek('kode = isi kolom Kode Unit pada tabel bagian 5', kode == kode_tabel,
        str([(a, b) for a, b in zip(kode, kode_tabel) if a != b][:3]))
    cek('judul unit terisi', all(u['judul_skema'] for u in hasil['units']))

    # ------------------------------------------------------------- peringatan
    print("\n3. PERINGATAN")
    cek('tidak ada peringatan pada dokumen ini', not hasil['peringatan'],
        str(hasil['peringatan']))

    # ------------------------------------------------- uji kasus duplikat/bad
    print("\n4. UJI SINTETIS: DUPLIKAT & KODE TIDAK VALID")
    tmp = project_root / 'output' / 'tmp_schema_test.docx'
    tmp.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.add_paragraph('Kemasan / Paket Kompetensi')
    t = d.add_table(rows=1, cols=3)
    for i, h in enumerate(('No', 'Kode Unit', 'Judul Unit')):
        t.rows[0].cells[i].text = h
    for no, kd, jd in [(1, 'J.63OPR00.009.2', 'Unit A'),
                       (2, 'J.63OPR00.009.2', 'Unit A lagi'),
                       (3, 'BUKAN-KODE', 'Unit B')]:
        row = t.add_row()
        row.cells[0].text, row.cells[1].text, row.cells[2].text = str(no), kd, jd
    d.add_paragraph('Persyaratan Dasar Pemohon Sertifikasi')
    d.save(tmp)

    uji = baca_skema(tmp)
    cek('duplikat terdeteksi',
        any('duplikat' in p.lower() for p in uji['peringatan']), str(uji['peringatan']))
    cek('kode tidak valid ditandai',
        any(not u['pola_kode_valid'] for u in uji['units'])
        and any('tidak sesuai pola' in p.lower() for p in uji['peringatan']))
    cek('kode tidak valid diberi catatan verifikasi',
        any(u['catatan'] for u in uji['units'] if not u['pola_kode_valid']))
    tmp.unlink(missing_ok=True)

    # ------------------------------------------------------------ keluaran
    print("\n5. BERKAS KELUARAN")
    data = json.loads(OUTPUT.read_text(encoding='utf-8'))
    cek('JSON tersimpan dan terbaca ulang', data['jumlah_unit'] == 10
        and [u['kode_unit'] for u in data['units']] == kode)

    print("\n" + "=" * 100)
    print(f"HASIL: {len(lulus)} lulus, {len(gagal)} gagal")
    if gagal:
        for g in gagal:
            print(f"  - {g}")
    print("=" * 100)
    return 0 if not gagal else 1


if __name__ == '__main__':
    sys.exit(main())
