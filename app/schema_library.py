"""
SKKNI Extractor - Perpustakaan dokumen skema sertifikasi (Fase 4C).

Menyimpan DOCX skema di dalam proyek, membaca daftar unit dengan
`app/schema_parser.py`, lalu mencocokkan kode unit ke perpustakaan PDF SKKNI
memakai `app/unit_matcher.py`.

Tidak ada parser baru untuk unit: seluruh pembacaan memakai modul Fase 3D.
"""

import hashlib
import json
import os
import re
import uuid
import zipfile
from datetime import datetime
from pathlib import Path

from app import library
from app.schema_parser import baca_skema
from app.unit_matcher import cocokkan, indeks_pdf

ROOT = Path(__file__).resolve().parent.parent
SKEMA_DIR = ROOT / 'data' / 'schema_library'
DOCX_STORE = SKEMA_DIR / 'docx'
META_DIR = SKEMA_DIR / 'documents'
TMP_DIR = SKEMA_DIR / 'tmp'

BATAS_UKURAN = 20 * 1024 * 1024        # 20 MB
BATAS_ENTRI_ZIP = 2000
BATAS_UNCOMPRESSED = 100 * 1024 * 1024
ID_POLA = re.compile(r'^[0-9a-f]{32}$')

POLA_JENIS = re.compile(r'jenis\s+skema\s*:\s*([^\n]+)', re.I)
POLA_NOMOR = re.compile(r'nomor\s+dokumen[ \t]*:[ \t]*([^\n\t]+)', re.I)

JUDUL_JSON = (
    ROOT / 'data' / 'json' / 'schema_content_creator',
    ROOT / 'data' / 'json' / 'cross_skkni',
    ROOT / 'data' / 'json' / 'browser_extraction',   # hasil ekstraksi dari browser
)


class GalatSkema(Exception):
    """Kesalahan aman untuk ditampilkan ke pengguna."""

    def __init__(self, pesan, kode=400):
        super().__init__(pesan)
        self.pesan = pesan
        self.kode = kode


# --------------------------------------------------------------------------
# Direktori & metadata
# --------------------------------------------------------------------------

def siapkan_direktori() -> None:
    for d in (DOCX_STORE, META_DIR, TMP_DIR):
        d.mkdir(parents=True, exist_ok=True)


def _tulis_json_atomik(path: Path, data: dict) -> None:
    sementara = path.with_suffix(path.suffix + '.tmp')
    with open(sementara, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.flush()
        os.fsync(f.fileno())
    os.replace(sementara, path)


def _baca_json(path: Path) -> dict | None:
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except Exception:
        return None


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for potong in iter(lambda: f.read(1024 * 1024), b''):
            h.update(potong)
    return h.hexdigest()


def daftar_skema() -> list[dict]:
    if not META_DIR.exists():
        return []
    hasil = []
    for p in META_DIR.glob('*.json'):
        d = _baca_json(p)
        if d and d.get('id'):
            hasil.append(d)
    hasil.sort(key=lambda d: d.get('tanggal_ditambah') or '', reverse=True)
    return hasil


def ambil_skema(schema_id: str) -> dict | None:
    if not ID_POLA.match(schema_id or ''):
        return None
    return _baca_json(META_DIR / f'{schema_id}.json')


def simpan_metadata(doc: dict) -> None:
    doc['diperbarui'] = datetime.now().isoformat(timespec='seconds')
    _tulis_json_atomik(META_DIR / f"{doc['id']}.json", doc)


def jalur_docx(doc: dict) -> Path:
    target = (ROOT / doc['nama_internal']).resolve()
    if ROOT.resolve() not in target.parents:
        raise GalatSkema('Lokasi dokumen tidak sah.')
    if not target.exists():
        raise GalatSkema('Berkas DOCX tidak ditemukan.', 404)
    return target


# --------------------------------------------------------------------------
# Validasi DOCX
# --------------------------------------------------------------------------

def validasi_docx(path: Path, cek_ekstensi: bool = True) -> dict:
    """Periksa signature ZIP, struktur DOCX, ukuran entri, dan makro."""
    if cek_ekstensi and path.suffix.lower() != '.docx':
        raise GalatSkema('Hanya dokumen .docx yang diterima.')

    with open(path, 'rb') as f:
        if f.read(4) != b'PK\x03\x04':
            raise GalatSkema('Berkas bukan DOCX yang sah (signature ZIP tidak ditemukan).')

    if not zipfile.is_zipfile(path):
        raise GalatSkema('Arsip DOCX rusak dan tidak dapat dibuka.')

    try:
        with zipfile.ZipFile(path) as z:
            nama = z.namelist()
            if '[Content_Types].xml' not in nama or 'word/document.xml' not in nama:
                raise GalatSkema('Struktur DOCX tidak lengkap (bagian Word tidak ditemukan).')
            if any(n.lower().endswith('vbaproject.bin') for n in nama):
                raise GalatSkema('Dokumen mengandung makro dan tidak diizinkan.')

            # Perlindungan ZIP bomb
            if len(nama) > BATAS_ENTRI_ZIP:
                raise GalatSkema('Arsip DOCX memiliki terlalu banyak berkas di dalamnya.')
            total = sum(i.file_size for i in z.infolist())
            if total > BATAS_UNCOMPRESSED:
                raise GalatSkema('Isi DOCX setelah dibuka melebihi batas aman.')
            if any(i.compress_size > 0 and i.file_size / i.compress_size > 200
                   for i in z.infolist()):
                raise GalatSkema('Arsip DOCX mencurigakan (rasio kompresi ekstrem).')
    except zipfile.BadZipFile:
        raise GalatSkema('Arsip DOCX rusak dan tidak dapat dibaca.')

    try:
        from docx import Document
        doc = Document(str(path))
        jumlah_paragraf = len(doc.paragraphs)
        jumlah_tabel = len(doc.tables)
    except Exception:
        raise GalatSkema('Dokumen DOCX tidak dapat dibaca (mungkin rusak atau terenkripsi).')

    return {'paragraf': jumlah_paragraf, 'tabel': jumlah_tabel}


# --------------------------------------------------------------------------
# Informasi skema
# --------------------------------------------------------------------------

def baca_info(path: Path) -> dict:
    """Ambil nama, jenis, dan nomor dokumen dari isi DOCX (tanpa mengarang)."""
    from docx import Document

    info = {'nama_skema': None, 'jenis_skema': None, 'nomor_dokumen': None,
            'judul_dokumen': None}
    try:
        doc = Document(str(path))
    except Exception:
        return info

    paragraf = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    teks = '\n'.join(paragraf)

    if paragraf and paragraf[0].upper().startswith('SKEMA SERTIFIKASI'):
        judul = paragraf[0]
        if len(paragraf) > 1 and len(paragraf[1]) <= 120:
            judul = f'{paragraf[0]} {paragraf[1]}'
        info['judul_dokumen'] = re.sub(r'\s+', ' ', judul).strip()
        # Nama skema: baris setelah judul, bila ada.
        if len(paragraf) > 1 and not paragraf[1].lower().startswith(('skema sertifikasi', 'nomor')):
            info['nama_skema'] = re.sub(r'\s+', ' ', paragraf[1]).strip()
        else:
            info['nama_skema'] = re.sub(r'\s+', ' ', paragraf[0]).strip()

    m = POLA_JENIS.search(teks)
    if m:
        info['jenis_skema'] = re.sub(r'\s+', ' ', m.group(1)).strip(' .')

    m = POLA_NOMOR.search(teks)
    if m:
        nilai = re.sub(r'\s+', ' ', m.group(1)).strip()
        if nilai and nilai != ':':
            info['nomor_dokumen'] = nilai
    return info


# --------------------------------------------------------------------------
# Sumber PDF & pencocokan
# --------------------------------------------------------------------------

def pdf_sumber() -> list[Path]:
    """
    PDF SKKNI yang menjadi sumber pencocokan:
    - semua PDF di data/pdf/
    - PDF unggahan yang terdaftar di metadata perpustakaan Fase 4B

    Berkas tanpa metadata (mis. sisa unggahan yang gagal) tidak dipakai agar
    tidak menimbulkan unit "hantu" pada hasil pencocokan.
    """
    hasil = list((ROOT / 'data' / 'pdf').glob('*.pdf'))
    dikenal = {d.get('nama_internal') for d in library.daftar_dokumen()}
    if library.PDF_STORE.exists():
        for p in library.PDF_STORE.glob('*.pdf'):
            if str(p.relative_to(ROOT)) in dikenal:
                hasil.append(p)
    return sorted(set(hasil))


_CACHE_INDEKS: dict = {'tanda': None, 'indeks': None}


def _tanda_sumber(paths: list[Path]) -> tuple:
    return tuple((str(p), p.stat().st_mtime, p.stat().st_size) for p in paths if p.exists())


def indeks_skkni(paksa: bool = False) -> dict:
    """Indeks unit PDF; di-cache selama berkas tidak berubah."""
    paths = pdf_sumber()
    tanda = _tanda_sumber(paths)
    if not paksa and _CACHE_INDEKS['tanda'] == tanda and _CACHE_INDEKS['indeks']:
        return _CACHE_INDEKS['indeks']
    idx = indeks_pdf(paths)
    _CACHE_INDEKS['tanda'] = tanda
    _CACHE_INDEKS['indeks'] = idx
    return idx


def _norm_judul(s: str) -> str:
    return re.sub(r'[^a-z0-9]+', '', (s or '').lower())


def cocokkan_units(units: list[dict], paksa_indeks: bool = False) -> dict:
    """Cocokkan unit skema ke perpustakaan; tandai beda judul."""
    idx = indeks_skkni(paksa=paksa_indeks)
    hasil = cocokkan(units, pdf_sumber(), indeks=idx)

    beda_judul = []
    for u in hasil['units']:
        # Istilah status Fase 4C: 'belum_ditemukan' (unit_matcher memakai
        # 'tidak_ditemukan'). Nilai tetap satu: tidak ada pencocokan.
        if u['status'] == 'tidak_ditemukan':
            u['status'] = 'belum_ditemukan'
        u['judul_berbeda'] = False
        if u.get('judul_pdf') and u.get('judul_skema'):
            if _norm_judul(u['judul_pdf']) != _norm_judul(u['judul_skema']):
                u['judul_berbeda'] = True
                u['catatan'].append('judul pada skema berbeda dengan judul pada PDF')
                beda_judul.append(u['kode_unit'])

    hasil['statistik']['judul_berbeda'] = len(beda_judul)
    hasil['judul_berbeda'] = beda_judul
    hasil['pdf_diperiksa'] = {k: v for k, v in idx['dokumen'].items()}
    hasil['diperiksa_pada'] = datetime.now().isoformat(timespec='seconds')
    return hasil


def simpan_pencocokan(doc: dict, hasil: dict) -> dict:
    doc['hasil_pencocokan'] = {
        'status': ('perlu_verifikasi'
                   if hasil['statistik']['tidak_ditemukan'] or hasil['statistik']['konflik']
                   else 'lengkap'),
        'diperiksa_pada': hasil['diperiksa_pada'],
        'ringkasan': hasil['statistik'],
        'pdf_diperiksa': hasil.get('pdf_diperiksa', {}),
        'units': {u['kode_unit']: u for u in hasil['units']},
        'tidak_ditemukan': hasil['tidak_ditemukan'],
        'konflik': [k['kode_unit'] if isinstance(k, dict) else k for k in hasil['konflik']],
        'judul_berbeda': hasil.get('judul_berbeda', []),
    }
    simpan_metadata(doc)
    return doc


# --------------------------------------------------------------------------
# Unggah
# --------------------------------------------------------------------------

def unggah(nama_asli: str, aliran, ukuran_dilaporkan: int | None = None) -> tuple[dict, bool]:
    siapkan_direktori()

    if ukuran_dilaporkan and ukuran_dilaporkan > BATAS_UKURAN:
        raise GalatSkema(f'Ukuran berkas melebihi batas {BATAS_UKURAN // (1024 * 1024)} MB.', 413)

    nama_bersih = Path(nama_asli or 'skema.docx').name
    if not nama_bersih.lower().endswith('.docx'):
        raise GalatSkema('Hanya dokumen .docx yang diterima (bukan .doc/.docm).')

    sid = uuid.uuid4().hex
    sementara = TMP_DIR / f'{sid}.part'
    h = hashlib.sha256()
    total = 0

    try:
        with open(sementara, 'wb') as f:
            while True:
                potong = aliran.read(1024 * 1024)
                if not potong:
                    break
                total += len(potong)
                if total > BATAS_UKURAN:
                    raise GalatSkema(
                        f'Ukuran berkas melebihi batas {BATAS_UKURAN // (1024 * 1024)} MB.', 413)
                h.update(potong)
                f.write(potong)
            f.flush()
            os.fsync(f.fileno())

        if total == 0:
            raise GalatSkema('Berkas kosong.')

        validasi_docx(sementara, cek_ekstensi=False)   # nama sementara berakhiran .part
        sha = h.hexdigest()

        for ada in daftar_skema():
            if ada.get('sha256') == sha:
                sementara.unlink(missing_ok=True)
                return ada, True

        tujuan = DOCX_STORE / f'{sid}.docx'
        os.replace(sementara, tujuan)

        info = baca_info(tujuan)
        doc = {
            'id': sid,
            'nama_asli': nama_bersih,
            'nama_internal': str(tujuan.relative_to(ROOT)),
            'sha256': sha,
            'ukuran_byte': total,
            'nama_skema': info['nama_skema'],
            'judul_dokumen': info['judul_dokumen'],
            'jenis_skema': info['jenis_skema'],
            'nomor_dokumen': info['nomor_dokumen'],
            'tanggal_ditambah': datetime.now().isoformat(timespec='seconds'),
            'status_identifikasi': 'perlu_verifikasi',
            'pesan_identifikasi': None,
            'jumlah_unit': 0,
            'units': [],
            'peringatan': [],
            'hasil_pencocokan': None,
            'verifikasi_manual': {},
        }

        try:
            skema = baca_skema(tujuan)
            doc['units'] = skema['units']
            doc['jumlah_unit'] = skema['jumlah_unit']
            doc['peringatan'] = skema['peringatan']
            doc['status_identifikasi'] = ('perlu_verifikasi' if skema['peringatan']
                                          or not skema['units'] else 'lengkap')
        except ValueError as e:
            doc['pesan_identifikasi'] = (
                f'Bagian daftar unit tidak dikenali: {e}. Periksa dokumen secara manual.')
        except Exception:
            doc['pesan_identifikasi'] = ('Dokumen tidak dapat dibaca sebagai skema '
                                         'sertifikasi. Periksa secara manual.')

        simpan_metadata(doc)

        if doc['units']:
            doc = simpan_pencocokan(doc, cocokkan_units(doc['units']))
        return doc, False
    finally:
        sementara.unlink(missing_ok=True)


# --------------------------------------------------------------------------
# Pemeriksaan ulang
# --------------------------------------------------------------------------

def cocokkan_ulang(doc: dict) -> dict:
    """Cocokkan ulang unit tersimpan dengan perpustakaan SKKNI saat ini."""
    if not doc.get('units'):
        raise GalatSkema('Dokumen ini tidak memiliki daftar unit untuk dicocokkan.', 409)

    # Segarkan informasi dokumen (nama/jenis/nomor) dari DOCX, bukan dari memori.
    try:
        info = baca_info(jalur_docx(doc))
        for k in ('nama_skema', 'judul_dokumen', 'jenis_skema', 'nomor_dokumen'):
            doc[k] = info.get(k)      # None berarti tidak ditemukan pada dokumen
    except GalatSkema:
        pass
    hasil = cocokkan_units(doc['units'], paksa_indeks=True)
    return simpan_pencocokan(doc, hasil)


# --------------------------------------------------------------------------
# Status unit (Fase 4C bagian G)
# --------------------------------------------------------------------------

def _json_unit(kode: str) -> dict | None:
    for folder in JUDUL_JSON:
        p = folder / f'{kode}.json'
        if p.exists():
            try:
                return json.loads(p.read_text(encoding='utf-8'))
            except Exception:
                return None
    return None


def peta_kelengkapan() -> dict[str, str]:
    """Peta kode unit -> status kelengkapan KUK dari metadata perpustakaan."""
    peta = {}
    for d in library.daftar_dokumen():
        for kode, unit in ((d.get('validasi_kelengkapan') or {}).get('unit') or {}).items():
            peta.setdefault(kode, unit.get('status'))
    return peta


def status_unit(doc: dict, kode: str, peta_kelengkapan: dict | None = None) -> dict:
    """Status berlapis untuk satu unit: pencocokan, ekstraksi, validasi, manual."""
    cocok = (doc.get('hasil_pencocokan') or {}).get('units', {}).get(kode) or {}
    data = _json_unit(kode)
    manual = (doc.get('verifikasi_manual') or {}).get(kode)
    peta = peta_kelengkapan if peta_kelengkapan is not None else globals()['peta_kelengkapan']()
    lengkap = peta.get(kode)

    return {
        'pencocokan': cocok.get('status', 'belum_ditemukan'),
        'teridentifikasi': cocok.get('status') == 'ditemukan',
        'sudah_diekstrak': bool(data),
        'valid_struktur': bool(data) and (data.get('validasi') or {}).get('status') == 'valid_struktur',
        'valid_kelengkapan': bool(data) and lengkap == 'valid_kelengkapan',
        'terverifikasi_manual': bool(manual),
        'judul_berbeda': bool(cocok.get('judul_berbeda')),
    }


def tandai_manual(doc: dict, kode_unit: str) -> dict:
    if not any(u['kode_unit'] == kode_unit for u in doc.get('units') or []):
        raise GalatSkema('Kode unit tidak ada pada skema ini.', 404)
    doc.setdefault('verifikasi_manual', {})[kode_unit] = {
        'waktu': datetime.now().isoformat(timespec='seconds'),
        'oleh': 'pengguna_lokal',
    }
    simpan_metadata(doc)
    return doc
